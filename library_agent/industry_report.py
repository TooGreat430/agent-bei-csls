"""Laporan "Price Competitiveness Analysis — B2B Segment" (one-pager industri).

Gaya mengikuti template PTPL: garis atas navy #003087, judul serif, tombol tahap Early/Next,
A. Price Index Highlight (kartu KPI), B. Price Gap Matrix (zona + segmen customer),
C. Executive Summary & Actionable Insights. HTML: satu file dengan tombol tahap.
PDF: satu halaman per tahap. PowerPoint: satu slide per tahap.
"""
from __future__ import annotations

import base64
import html as _html
import io
import json
import logging
import random
from typing import Any

from . import industry_data as idm

logger = logging.getLogger(__name__)
C = {"navy": "003087", "navy_mid": "1a4fa0", "navy_light": "d6e4f5", "text": "1a1a1a", "muted": "5c6878",
     "border": "ccd4dc", "border_light": "e8ecf0", "bg_alt": "f5f7fa", "red": "c0392b", "red_bg": "fdf0ef",
     "amber": "d4720a", "amber_bg": "fdf5e6", "green": "1a7a3c", "green_bg": "edf7f1"}
TONE_FG = {"red": C["red"], "amber": C["amber"], "green": C["green"], "na": C["muted"]}
TONE_BG = {"red": C["red_bg"], "amber": C["amber_bg"], "green": C["green_bg"], "na": "FFFFFF"}


def pct(v: float | None) -> str:
    return "—" if v is None else f"{'+' if v >= 0 else ''}{v:.1f}%"


def idr(v: float | None) -> str:
    return "—" if v is None else f"{round(v):,}".replace(",", ".")


def delta(cur: float | None, prev: float | None) -> tuple[str, str]:
    if cur is None or prev is None:
        return "", "flat"
    d = round(cur - prev, 1)
    if abs(d) < 0.05:
        return "= 0%", "flat"
    return f"{'↑' if d > 0 else '↓'} {'+' if d > 0 else ''}{d}%", "up" if d > 0 else "dn"


def _e(t: Any) -> str:
    return _html.escape("" if t is None else str(t))


# ==========================================================================
# Kartu KPI & narasi
# ==========================================================================
def kpi_cards(stage: dict, competitor: str) -> list[dict]:
    kpi = stage["kpi"]
    cards = []
    seg = kpi.get("segment") or {}
    if seg.get("name"):
        d, cls = delta(seg["current"], seg["prev"])
        cards.append({"dim": "CUSTOMER SEGMENT", "name": seg["name"], "value": seg["current"], "delta": d, "dcls": cls,
                      "tag": None, "tone": idm.tone(seg["current"]),
                      "text": f"Segmen dengan gap terendah pada stage ini: {pct(seg['current'])}. Basis = rata-rata produk "
                              f"fokus yang tersedia (n={seg['n']} product cells)."})
    cats = [(k, v) for k, v in kpi.items() if k != "segment" and v.get("current") is not None]
    min_cat = min(cats, key=lambda kv: kv[1]["current"])[0] if cats else None
    for cat, v in cats:
        cur = v["current"]
        tag = ("⚠ MOST NEGATIVE" if cat == min_cat and cur < 0 else "TOP PERFORMER" if cur > 10 else
               "PERLU DIPANTAU" if cur >= 0 else "TIDAK KOMPETITIF")
        best, worst = v.get("best_zone"), v.get("worst_zone")
        za = v.get("zone_avg") or {}
        if cur < 0:
            text = f"Kategori dengan gap negatif — rata-rata {competitor.title()} lebih murah."
        elif cur > 10:
            text = "Kategori kompetitif secara nasional."
        else:
            text = "Kategori dengan selisih tipis (0–10%)."
        if best:
            text += f" {best} terkuat ({pct(za.get(best))})."
        if worst and worst != best:
            text += f" Titik lemah: {worst} ({pct(za.get(worst))})."
        d, cls = delta(cur, v.get("prev"))
        cards.append({"dim": cat.upper(), "name": v.get("label", cat), "value": cur, "delta": d, "dcls": cls,
                      "tag": tag, "tone": idm.tone(cur), "text": text})
    return cards[:6]


NARRATIVE_SCHEMA = {
    "type": "object",
    "properties": {"stages": {"type": "array", "items": {"type": "object", "properties": {
        "stage": {"type": "string"},
        "ringkasan": {"type": "array", "items": {"type": "object", "properties": {
            "judul": {"type": "string"}, "isi": {"type": "string"}}, "required": ["judul", "isi"]}},
        "aksi": {"type": "array", "items": {"type": "object", "properties": {
            "label": {"type": "string", "enum": ["URGENT", "LEVERAGE", "MONITOR"]},
            "judul": {"type": "string"}, "isi": {"type": "string"}}, "required": ["label", "judul", "isi"]}}},
        "required": ["stage", "ringkasan", "aksi"]}}},
    "required": ["stages"],
}


def _compact(ds: dict) -> dict:
    out = {"periode": ds["period_a"]["label"], "pembanding": ds["period_b"]["label"] if ds["period_b"] else None, "stages": {}}
    for st, data in ds["stages"].items():
        cells = []
        for prod, row in data["table1"].items():
            for z, c in row.items():
                if c and c["current"] is not None:
                    cells.append({"produk": idm.short_name(prod), "zona": z, "gap": pct(c["current"]), "sebelumnya": pct(c["prev"])})
        segs = []
        for ch, row in data["table2"].items():
            for prod, c in row.items():
                if c and c["current"] is not None:
                    segs.append({"segmen": ch, "produk": idm.short_name(prod), "gap": pct(c["current"])})
        kpi = {k: {"gap": pct(v.get("current")), "sebelumnya": pct(v.get("prev")), **({"nama": v.get("name")} if k == "segment" else {})}
               for k, v in data["kpi"].items()}
        out["stages"][st] = {"kpi": kpi, "zona": cells, "segmen": segs}
    return out


def narrative_prompt(ds: dict, competitor: str) -> str:
    return f"""Anda analis Market Research & Intelligence PTPL untuk segmen B2B/industri. Tulis narasi one-pager
"Price Competitiveness Analysis — B2B Segment" dari DATA berikut.

ATURAN
- Gap (%) = (harga {competitor} − HTD PTPL+3%) / HTD PTPL+3%. POSITIF = PTPL kompetitif; NEGATIF = {competitor}
  lebih murah. Lampu: > 10% kompetitif, 0–10% hati-hati, < 0% tidak kompetitif.
- HANYA pakai angka yang ada di DATA, ditulis persis (mis. "+15.8%", "-16.2%"). Jangan menghitung angka baru.
- Untuk setiap stage di DATA: "ringkasan" 3-5 poin (judul singkat + isi 1 kalimat), "aksi" 2-3 poin dengan
  label URGENT (gap negatif yang perlu tindakan), LEVERAGE (keunggulan yang bisa dimanfaatkan), atau MONITOR.
- Bahasa Indonesia ringkas untuk manajemen. Jangan menyebut nama tabel atau filter.

DATA
{json.dumps(_compact(ds), ensure_ascii=False)}
"""


def generate_narrative(ds: dict, competitor: str) -> dict[str, dict]:
    from .price_dashboard import ungrounded_numbers
    from .report_engine import _generate_json

    try:
        raw = _generate_json(narrative_prompt(ds, competitor), NARRATIVE_SCHEMA)
    except Exception:  # noqa: BLE001
        logger.exception("Narasi industri gagal dibuat; laporan tetap dibuat tanpa narasi")
        return {}
    nums = idm.dataset_numbers(ds)
    out = {}
    for item in raw.get("stages") or []:
        st = str(item.get("stage", "")).upper()
        st = "EARLY" if "EARLY" in st else "NEXT" if "NEXT" in st else st
        ok = lambda t: not ungrounded_numbers(t, nums)  # noqa: E731
        out[st] = {"ringkasan": [r for r in item.get("ringkasan") or [] if ok(f"{r.get('judul')} {r.get('isi')}")],
                   "aksi": [a for a in item.get("aksi") or [] if ok(f"{a.get('judul')} {a.get('isi')}")]}
    return out


# ==========================================================================
# HTML
# ==========================================================================
CSS = f"""
@import url('https://fonts.googleapis.com/css2?family=Source+Serif+4:wght@300;400;600&family=Inter:wght@300;400;500;600&display=swap');
*{{box-sizing:border-box;margin:0;padding:0}}
body{{font-family:'Inter',Arial,sans-serif;background:#fff;color:#{C['text']};font-size:10.5px;line-height:1.45;padding:24px 28px 20px;min-width:1160px;max-width:1360px;margin:0 auto}}
.serif{{font-family:'Source Serif 4',Georgia,serif}}
.header{{border-top:3px solid #{C['navy']};padding:5px 0 10px;display:flex;justify-content:space-between;align-items:flex-start;margin-bottom:16px;border-bottom:.5px solid #{C['border']}}}
.brand{{display:flex;align-items:center;gap:15px}}.brand img{{height:40px}}
.doc-title{{font-family:'Source Serif 4',Georgia,serif;font-size:18px;font-weight:600;color:#{C['navy']};line-height:1.2;margin-bottom:2px}}
.doc-sub{{font-size:10px;color:#{C['muted']}}}.meta{{text-align:right}}.meta-line{{font-size:8px;color:#{C['muted']};line-height:1.7}}
.meta-period{{display:inline-block;background:#{C['navy']};color:#fff;font-size:9px;font-weight:500;padding:2px 8px;letter-spacing:.5px;margin-bottom:4px}}
.stage-filter{{display:flex;align-items:center;gap:8px;margin-bottom:14px}}
.sf-label{{font-size:9px;color:#{C['muted']};font-weight:600;text-transform:uppercase;letter-spacing:.6px;margin-right:4px}}
.stage-btn{{font-family:'Inter',Arial,sans-serif;font-size:9.5px;font-weight:600;letter-spacing:.4px;padding:5px 14px;border:1px solid #{C['navy']};background:#fff;color:#{C['navy']};cursor:pointer;text-transform:uppercase}}
.stage-btn.active{{background:#{C['navy']};color:#fff}}.stage-hint{{font-size:8.5px;color:#{C['muted']};margin-left:6px}}
.section-head{{display:flex;align-items:center;gap:5px;margin-bottom:10px}}.bar{{width:3px;height:14px;background:#{C['navy']}}}
.sh-label{{font-size:9.5px;font-weight:600;text-transform:uppercase;letter-spacing:1.2px;color:#{C['navy']}}}.rule{{flex:1;height:.5px;background:#{C['border']}}}
.sh-note{{font-size:9px;color:#{C['muted']}}}
.legend{{font-size:8.5px;color:#{C['muted']};margin:-4px 0 8px}}.legend span{{display:inline-block;width:10px;height:10px;vertical-align:middle;margin:0 3px 0 8px;border:1px solid}}
.kpi-grid{{display:grid;grid-template-columns:repeat(6,1fr);gap:8px;margin-bottom:8px}}
.kpi-card{{border:.5px solid #{C['border']};padding:10px 10px 8px;position:relative}}
.kpi-card::before{{content:'';position:absolute;top:0;left:0;right:0;height:3px}}
.k-red::before{{background:#{C['red']}}}.k-amber::before{{background:#{C['amber']}}}.k-green::before{{background:#{C['green']}}}.k-na::before{{background:#{C['border']}}}
.kpi-dim{{font-size:8.5px;font-weight:600;text-transform:uppercase;letter-spacing:.8px;color:#{C['muted']};margin-bottom:3px}}
.kpi-name{{font-family:'Source Serif 4',Georgia,serif;font-size:12px;font-weight:600;margin-bottom:7px}}
.kpi-row{{display:flex;align-items:baseline;gap:8px;padding-bottom:6px;border-bottom:.5px solid #{C['border_light']}}}
.kpi-num{{font-size:22px;font-weight:300;font-family:'Source Serif 4',Georgia,serif;line-height:1}}
.c-red{{color:#{C['red']}}}.c-amber{{color:#{C['amber']}}}.c-green{{color:#{C['green']}}}.c-na{{color:#{C['muted']}}}
.kpi-trend{{font-size:9px;font-weight:500}}.up{{color:#{C['green']}}}.dn{{color:#{C['red']}}}.flat{{color:#{C['muted']}}}
.kpi-text{{font-size:9px;color:#{C['muted']};line-height:1.55;margin-top:6px}}
.tag{{display:inline-block;font-size:8px;font-weight:600;letter-spacing:.5px;text-transform:uppercase;padding:1px 5px;margin-bottom:4px}}
.tag.best{{background:#{C['green_bg']};color:#{C['green']};border:.5px solid #a0cfb2}}.tag.warn{{background:#{C['red_bg']};color:#{C['red']};border:.5px solid #e8b0aa}}
.tag.mid{{background:#{C['amber_bg']};color:#{C['amber']};border:.5px solid #eccb94}}
.kpi-note{{font-size:8.5px;color:#{C['muted']};margin-bottom:16px}}
.table-block{{border:.5px solid #{C['border']};margin-bottom:16px}}
table{{width:100%;border-collapse:collapse}}
.grp th{{font-size:9px;padding:6px 10px;color:#fff;text-transform:uppercase;letter-spacing:.4px}}
.grp .gz{{background:#{C['navy']}}}.grp .gs{{background:#{C['amber_bg']};color:#{C['amber']}}}
thead tr.cols{{background:#{C['bg_alt']};border-bottom:1px solid #{C['border']}}}
thead th{{font-size:8.5px;font-weight:600;text-transform:uppercase;color:#{C['navy']};padding:5px 6px;text-align:center;white-space:nowrap;border-right:.5px solid #{C['border_light']}}}
thead th:first-child{{text-align:left}}th small{{font-weight:400;color:#{C['muted']}}}
tr.cat td{{text-align:left;background:#{C['bg_alt']};font-size:9px;font-weight:600;color:#{C['navy']};letter-spacing:.6px;padding:4px 8px;text-transform:uppercase}}
td{{border-top:.5px solid #{C['border_light']};border-right:.5px solid #{C['border_light']};padding:4px 4px;text-align:center;vertical-align:middle}}
td.prod{{text-align:left;padding-left:8px;font-size:10px;white-space:nowrap}}
.cv{{display:inline-block;font-family:'Courier New',monospace;font-size:10px;font-weight:600;padding:1px 5px}}
.cv-red{{background:#{C['red_bg']};color:#{C['red']}}}.cv-amber{{background:#{C['amber_bg']};color:#{C['amber']}}}.cv-green{{background:#{C['green_bg']};color:#{C['green']}}}.cv-na{{color:#{C['muted']}}}
.sub{{font-size:8px}}.chip{{display:inline-block;font-size:7.5px;background:#{C['navy_light']};color:#{C['navy']};padding:0 4px;margin-top:1px}}
.mm{{font-size:7px;color:#{C['muted']};white-space:nowrap}}.mm b{{color:#{C['text']}}}
.sec-c{{display:grid;grid-template-columns:1fr 1fr;gap:12px}}.box{{border:.5px solid #{C['border']};padding:12px 14px}}
.box h3{{font-family:'Source Serif 4',Georgia,serif;font-size:12px;color:#{C['navy']};text-transform:uppercase;letter-spacing:.6px;margin-bottom:8px}}
.box ul{{padding-left:14px}}.box li{{font-size:9.5px;margin-bottom:7px;border-bottom:.5px dashed #{C['border_light']};padding-bottom:6px}}.box li b{{color:#{C['navy']}}}
.act{{display:flex;gap:8px;align-items:flex-start;font-size:9.5px;margin-bottom:8px;border-bottom:.5px dashed #{C['border_light']};padding-bottom:7px}}
.lab{{font-size:8px;font-weight:700;color:#fff;padding:2px 6px;white-space:nowrap}}.lab.URGENT{{background:#{C['red']}}}.lab.LEVERAGE{{background:#{C['navy']}}}.lab.MONITOR{{background:#{C['amber']}}}
.footer{{border-top:.5px solid #{C['border']};margin-top:10px;padding-top:8px;font-size:8px;color:#{C['muted']};display:flex;justify-content:space-between;gap:20px}}
.stage{{display:none}}.stage.on{{display:block}}.empty{{font-size:10px;color:#{C['muted']};padding:16px}}
"""


def _logo_uri() -> str:
    from .dashboard_render import logo_bytes

    lb = logo_bytes()
    return f"data:image/png;base64,{base64.b64encode(lb).decode()}" if lb else ""


def _cell_html(cell: dict | None, prev_label: str) -> str:
    if not cell or cell.get("current") is None:
        if cell and cell.get("prev") is not None:
            return f'<span class="cv cv-na">—</span><div class="sub flat">{_e(prev_label)}: {pct(cell["prev"])}</div>'
        return '<span class="cv cv-na">—</span>'
    t = idm.tone(cell["current"])
    d, cls = delta(cell["current"], cell.get("prev"))
    brands = ", ".join(b.title() for b in cell.get("brands") or [])
    return (f'<span class="cv cv-{t}">{pct(cell["current"])}</span>'
            + (f'<div class="sub {cls}">{_e(d)}</div>' if d else "")
            + (f'<div><span class="chip">{_e(brands)}</span></div>' if brands else "")
            + f'<div class="mm">Max <b>{idr(cell["max"])}</b> Avg <b>{idr(cell["avg"])}</b> Min <b>{idr(cell["min"])}</b></div>')


def _stage_html(st: str, data: dict, narr: dict, ds: dict, competitor: str) -> str:
    pa, pb = ds["period_a"], ds["period_b"]
    prev_label = pb["short"] if pb else ""
    if not data["has_data"]:
        return f'<div class="empty">Tidak ada data {idm.STAGE_LABEL[st]} untuk periode {pa["label"]}.</div>'
    cards = ""
    for c in kpi_cards(data, competitor):
        tag = ""
        if c["tag"]:
            cls = "warn" if "NEGATIVE" in c["tag"] or "TIDAK" in c["tag"] else "best" if "TOP" in c["tag"] else "mid"
            tag = f'<span class="tag {cls}">{_e(c["tag"])}</span>'
        cards += (f'<div class="kpi-card k-{c["tone"]}"><div class="kpi-dim">{_e(c["dim"])}</div>{tag}'
                  f'<div class="kpi-name">{_e(c["name"])}</div><div class="kpi-row"><span class="kpi-num c-{c["tone"]}">'
                  f'{pct(c["value"])}</span><span class="kpi-trend {c["dcls"]}">{_e(c["delta"])}</span></div>'
                  f'<div class="kpi-text">{_e(c["text"])}</div></div>')
    nseg = len(ds["segments"])
    head = ('<thead><tr class="grp"><th></th><th colspan="3" class="gz">Gap harga per zona wilayah vs produk fokus</th>'
            f'<th colspan="{nseg}" class="gs">Gap harga per segment customer vs produk fokus</th></tr>'
            '<tr class="cols"><th>Kategori / Produk</th>'
            + "".join(f'<th>{z}<br><small>{idm.ZONE_SUB[z]}</small></th>' for z in idm.ZONES)
            + "".join(f"<th>{_e(c)}</th>" for c in ds["segments"]) + "</tr></thead>")
    body = ""
    for cat in [c for c in idm.CATEGORY_ORDER if any(f["category"] == c for f in ds["focus"])] + \
            sorted({f["category"] for f in ds["focus"]} - set(idm.CATEGORY_ORDER)):
        body += f'<tr class="cat"><td colspan="{4 + len(ds["segments"])}">{_e(cat)}</td></tr>'
        for f in [f for f in ds["focus"] if f["category"] == cat]:
            p = f["name"]
            body += (f'<tr><td class="prod">{_e(idm.short_name(p))}</td>'
                     + "".join(f"<td>{_cell_html(data['table1'][p][z], prev_label)}</td>" for z in idm.ZONES)
                     + "".join(f"<td>{_cell_html(data['table2'][c][p], prev_label)}</td>" for c in ds["segments"]) + "</tr>")
    n = narr.get(st) or {}
    summ = "".join(f"<li><b>{_e(r.get('judul'))}</b> — {_e(r.get('isi'))}</li>" for r in n.get("ringkasan", [])) or \
        "<li>Narasi belum tersedia.</li>"
    acts = "".join(f'<div class="act"><span class="lab {_e(a.get("label"))}">{_e(a.get("label"))}</span><div><b>{_e(a.get("judul"))}</b> — '
                   f'{_e(a.get("isi"))}</div></div>' for a in n.get("aksi", []))
    trend = f"Trend = Δ vs {pb['label']}" if pb else "Tanpa pembanding"
    return (f'<div class="section-head"><div class="bar"></div><div class="sh-label">A &nbsp; Price Index Highlight</div>'
            f'<div class="rule"></div><div class="sh-note">Current = {_e(pa["label"])} | {_e(trend)}</div></div>'
            f'<div class="legend">Traffic light:<span style="border-color:#{C["green"]}"></span>Gap &gt;10% (kompetitif)'
            f'<span style="border-color:#{C["amber"]}"></span>Gap 0–10% (hati-hati)<span style="border-color:#{C["red"]}"></span>Gap &lt;0% (tidak kompetitif)</div>'
            f'<div class="kpi-grid">{cards}</div><div class="kpi-note">KPI kategori = simple average product × zona yang tersedia; '
            f'KPI segment = simple average product × segment yang tersedia. Gap zona memakai HTD+3% pada baris HTD yang sama.</div>'
            f'<div class="section-head"><div class="bar"></div><div class="sh-label">B &nbsp; Price Gap Matrix — Produk Fokus</div>'
            f'<div class="rule"></div><div class="sh-note">Current = Gap {_e(pa["label"].split()[0])} vs HTD PTPL+3% | Sub = Δ vs {_e((pb or {}).get("label", "-").split()[0])} | '
            f'Angka = Max / Avg / Min Harga {_e(competitor.title())}</div></div>'
            f'<div class="table-block"><table>{head}<tbody>{body}</tbody></table></div>'
            f'<div class="section-head"><div class="bar"></div><div class="sh-label">C &nbsp; Executive Summary &amp; Actionable Insights</div>'
            f'<div class="rule"></div></div><div class="sec-c"><div class="box"><h3>Executive Summary</h3><ul>{summ}</ul></div>'
            f'<div class="box"><h3>Actionable Insights</h3>{acts or "<div class=act>Belum ada rekomendasi.</div>"}</div></div>')


def render_html(ds: dict, narr: dict, meta: dict) -> bytes:
    pa, pb = ds["period_a"], ds["period_b"]
    comp = meta["competitor"]
    stages = "".join(f'<div class="stage{" on" if i == 0 else ""}" id="st-{st}">{_stage_html(st, ds["stages"][st], narr, ds, comp)}</div>'
                     for i, st in enumerate(idm.STAGES))
    buttons = "".join(f'<button class="stage-btn{" active" if i == 0 else ""}" data-s="{st}">{idm.STAGE_LABEL[st]}</button>'
                      for i, st in enumerate(idm.STAGES))
    n_cat = len({f["category"] for f in ds["focus"]})
    period_line = f"{pb['label'].split()[0]} – {pa['label']}" if pb else pa["label"]
    script = ("const H={EARLY:'%s',NEXT:'%s'};document.querySelectorAll('.stage-btn').forEach(b=>b.onclick=()=>{"
              "document.querySelectorAll('.stage-btn').forEach(x=>x.classList.remove('active'));"
              "document.querySelectorAll('.stage').forEach(x=>x.classList.remove('on'));b.classList.add('active');"
              "document.getElementById('st-'+b.dataset.s).classList.add('on');document.getElementById('hint').textContent=H[b.dataset.s];})"
              % (idm.STAGE_HINT["EARLY"], idm.STAGE_HINT["NEXT"]))
    doc = (f"<!DOCTYPE html><html lang='id'><head><meta charset='utf-8'><title>{_e(meta['report_title'])}</title><style>{CSS}</style></head><body>"
           f"<div class='header'><div class='brand'><img src='{_logo_uri()}' alt='Pertamina Lubricants'/><div>"
           f"<div class='doc-title'>Price Competitiveness Analysis — B2B Segment</div>"
           f"<div class='doc-sub'>{_e(meta['company'])} · {len(ds['segments'])} Segmen Customer Fokus · 3 Zona Wilayah · {n_cat} Kategori Produk</div></div></div>"
           f"<div class='meta'><span class='meta-period'>{_e(pa['label'])} (Most Recent)</span><div class='meta-line'>Periode: <strong>{_e(period_line)}</strong></div>"
           f"<div class='meta-line'>Gap positif = PTPL kompetitif | Gap negatif = {_e(comp.title())} lebih murah</div></div></div>"
           f"<div class='stage-filter'><span class='sf-label'>Main Stage:</span>{buttons}<span class='stage-hint' id='hint'>{idm.STAGE_HINT['EARLY']}</span></div>"
           f"{stages}<div class='footer'><div>Gap = (AVG Harga {_e(comp.title())} − HTD PTPL+3%) / HTD PTPL+3% · Zone gap uses row-matched HTD+3%; "
           f"Segment gap = average of available row-level gaps · Data Source: {_e(ds.get('source_label', idm.SOURCE_LABEL))}, {_e(pa['label'])} · "
           f"Confidential — Internal Use Only</div><div class='serif' style='color:#{C['navy']};font-weight:600'>{_e(meta['company'])}</div></div>"
           f"<script>{script}</script></body></html>")
    return doc.encode("utf-8")


# ==========================================================================
# PDF (satu halaman per tahap)
# ==========================================================================
def render_pdf(ds: dict, narr: dict, meta: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph, Table, TableStyle

    from .dashboard_render import logo_bytes, plain
    from .report_render import _register_fonts

    regular, bold, _ = _register_fonts()
    W, H = 1000, 720
    buf = io.BytesIO()
    cv = canvas.Canvas(buf, pagesize=(W, H))
    cv.setTitle(meta["report_title"])
    hx = lambda c: colors.HexColor("#" + c)  # noqa: E731
    lb = logo_bytes()
    comp = meta["competitor"].title()
    pa, pb = ds["period_a"], ds["period_b"]

    def para(text, size=7, color=C["text"], font=None, align=0, markup=False, leading=None):
        st = ParagraphStyle("p", fontName=font or regular, fontSize=size, leading=leading or size * 1.3, textColor=hx(color), alignment=align)
        return Paragraph(plain(text) if markup else _e(plain(text)), st)

    def section(label, note, y):
        cv.setFillColor(hx(C["navy"]))
        cv.rect(24, y - 9, 2.5, 10, fill=1, stroke=0)
        cv.setFont(bold, 7.5)
        cv.drawString(31, y - 7, label.upper())
        if note:
            cv.setFont(regular, 6.3)
            cv.setFillColor(hx(C["muted"]))
            cv.drawRightString(W - 24, y - 7, plain(note))
        cv.setStrokeColor(hx(C["border"]))
        cv.setLineWidth(.4)
        cv.line(31 + cv.stringWidth(label.upper(), bold, 7.5) + 6, y - 5, W - 24 - (cv.stringWidth(plain(note), regular, 6.3) + 6 if note else 0), y - 5)
        return y - 16

    for st in idm.STAGES:
        data = ds["stages"][st]
        cv.setStrokeColor(hx(C["navy"]))
        cv.setLineWidth(2.2)
        cv.line(24, H - 20, W - 24, H - 20)
        if lb:
            cv.drawImage(ImageReader(io.BytesIO(lb)), 24, H - 58, width=96, height=31, mask="auto")
        cv.setFillColor(hx(C["navy"]))
        cv.setFont("Times-Bold", 14)
        cv.drawString(130, H - 40, "Price Competitiveness Analysis — B2B Segment")
        cv.setFont(regular, 7.5)
        cv.setFillColor(hx(C["muted"]))
        cv.drawString(130, H - 52, f"{meta['company']} · {len(ds['segments'])} Segmen Customer Fokus · 3 Zona Wilayah · "
                                   f"{len({f['category'] for f in ds['focus']})} Kategori Produk")
        tw = cv.stringWidth(f"{pa['label']} (Most Recent)", bold, 7) + 10
        cv.setFillColor(hx(C["navy"]))
        cv.rect(W - 24 - tw, H - 38, tw, 11, fill=1, stroke=0)
        cv.setFillColor(colors.white)
        cv.setFont(bold, 7)
        cv.drawRightString(W - 29, H - 35, f"{pa['label']} (Most Recent)")
        cv.setFillColor(hx(C["muted"]))
        cv.setFont(regular, 6.3)
        cv.drawRightString(W - 24, H - 47, f"Main stage: {idm.STAGE_LABEL[st]} · Periode: "
                                           f"{(pb['label'].split()[0] + ' – ') if pb else ''}{pa['label']}")
        cv.drawRightString(W - 24, H - 56, f"Gap positif = PTPL kompetitif | Gap negatif = {comp} lebih murah")
        cv.setStrokeColor(hx(C["border"]))
        cv.setLineWidth(.4)
        cv.line(24, H - 64, W - 24, H - 64)
        y = H - 74
        if not data["has_data"]:
            cv.setFont(regular, 9)
            cv.drawString(24, y - 20, f"Tidak ada data {idm.STAGE_LABEL[st]} untuk periode {pa['label']}.")
            cv.showPage()
            continue
        y = section("A  Price Index Highlight", f"Current = {pa['label']}" + (f" | Trend = Δ vs {pb['label']}" if pb else ""), y)
        cards = kpi_cards(data, meta["competitor"])
        cw = (W - 48 - 5 * 6) / 6
        heights = []
        for i, c in enumerate(cards):
            flows = [para(c["dim"], 5.8, C["muted"], bold)]
            if c["tag"]:
                flows.append(para(c["tag"], 5.5, C["red"] if "NEG" in c["tag"] or "TIDAK" in c["tag"] else C["green"] if "TOP" in c["tag"] else C["amber"], bold))
            flows += [para(c["name"], 8, C["text"], bold),
                      para(f'<font size="14" color="#{TONE_FG[c["tone"]]}">{pct(c["value"])}</font>  '
                           f'<font size="6.5" color="#{C["green"] if c["dcls"] == "up" else C["red"] if c["dcls"] == "dn" else C["muted"]}">{_e(c["delta"])}</font>',
                           7, markup=True, leading=18),
                      para(c["text"], 5.8, C["muted"])]
            t = Table([[flows]], colWidths=[cw])
            t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .4, hx(C["border"])), ("LINEABOVE", (0, 0), (-1, 0), 2.5, hx(TONE_FG[c["tone"]])),
                                   ("TOPPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 5)]))
            t.wrap(cw, 200)
            heights.append(t._height)
            t.drawOn(cv, 24 + i * (cw + 6), y - t._height)
        y -= max(heights or [0]) + 12
        y = section("B  Price Gap Matrix — Produk Fokus",
                    f"Current = Gap {pa['label'].split()[0]} vs HTD PTPL+3% | Sub = Δ vs {pb['label'].split()[0] if pb else '-'} | Angka = Max/Avg/Min Harga {comp}", y)
        prev_label = pb["short"] if pb else ""
        head1 = [para("", 6)] + [para("GAP HARGA PER ZONA WILAYAH VS PRODUK FOKUS", 6, "FFFFFF", bold, 1), "", ""] + \
                [para("GAP HARGA PER SEGMENT CUSTOMER VS PRODUK FOKUS", 6, C["amber"], bold, 1)] + [""] * (len(ds["segments"]) - 1)
        head2 = [para("KATEGORI / PRODUK", 6, C["navy"], bold)] + \
                [para(f"{z.upper()}<br/>{idm.ZONE_SUB[z]}", 5.8, C["navy"], bold, 1, markup=True) for z in idm.ZONES] + \
                [para(c.upper(), 5.8, C["navy"], bold, 1) for c in ds["segments"]]
        last = 3 + len(ds["segments"])
        rows, cmds = [head1, head2], [("SPAN", (1, 0), (3, 0)), ("SPAN", (4, 0), (last, 0)), ("BACKGROUND", (1, 0), (3, 0), hx(C["navy"])),
                                      ("BACKGROUND", (4, 0), (last, 0), hx(C["amber_bg"])), ("BACKGROUND", (0, 1), (-1, 1), hx(C["bg_alt"])),
                                      ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("GRID", (0, 1), (-1, -1), .3, hx(C["border_light"])),
                                      ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]

        def cellp(cell):
            if not cell or cell.get("current") is None:
                return para("—", 7, C["muted"], align=1)
            t = idm.tone(cell["current"])
            d, cls = delta(cell["current"], cell.get("prev"))
            brands = ", ".join(b.title() for b in cell.get("brands") or [])
            txt = (f'<font name="{bold}" color="#{TONE_FG[t]}">{pct(cell["current"])}</font>'
                   + (f'<br/><font size="5.5" color="#{C["green"] if cls == "up" else C["red"] if cls == "dn" else C["muted"]}">{_e(d)}</font>' if d else "")
                   + (f'<br/><font size="5" color="#{C["navy"]}">{_e(brands)}</font>' if brands else "")
                   + f'<br/><font size="4.6" color="#{C["muted"]}">Max {idr(cell["max"])} Avg {idr(cell["avg"])} Min {idr(cell["min"])}</font>')
            return para(txt, 7, align=1, markup=True)

        ri = 2
        for cat in [c for c in idm.CATEGORY_ORDER if any(f["category"] == c for f in ds["focus"])] + \
                sorted({f["category"] for f in ds["focus"]} - set(idm.CATEGORY_ORDER)):
            rows.append([para(cat.upper(), 6.2, C["navy"], bold)] + [""] * (last))
            cmds += [("SPAN", (0, ri), (-1, ri)), ("BACKGROUND", (0, ri), (-1, ri), hx(C["bg_alt"]))]
            ri += 1
            for f in [f for f in ds["focus"] if f["category"] == cat]:
                p = f["name"]
                line = [para(idm.short_name(p), 7)] + [cellp(data["table1"][p][z]) for z in idm.ZONES] + \
                       [cellp(data["table2"][c][p]) for c in ds["segments"]]
                for ci, cell in enumerate([data["table1"][p][z] for z in idm.ZONES] + [data["table2"][c][p] for c in ds["segments"]], start=1):
                    if cell and cell.get("current") is not None:
                        cmds.append(("BACKGROUND", (ci, ri), (ci, ri), hx(TONE_BG[idm.tone(cell["current"])])))
                rows.append(line)
                ri += 1
        first = 92
        t = Table(rows, colWidths=[first] + [(W - 48 - first) / last] * last)
        t.setStyle(TableStyle(cmds))
        t.wrap(W - 48, H)
        t.drawOn(cv, 24, y - t._height)
        y -= t._height + 12
        y = section("C  Executive Summary & Actionable Insights", "", y)
        n = narr.get(st) or {}
        half = (W - 48 - 12) / 2
        s_flows = [para("EXECUTIVE SUMMARY", 7.5, C["navy"], bold)] + \
                  [para(f"• <b>{_e(r.get('judul'))}</b> — {_e(r.get('isi'))}", 6.6, markup=True) for r in n.get("ringkasan", [])]
        a_flows = [para("ACTIONABLE INSIGHTS", 7.5, C["navy"], bold)] + \
                  [para(f'<font color="#{C["red"] if a.get("label") == "URGENT" else C["navy"] if a.get("label") == "LEVERAGE" else C["amber"]}">'
                        f'<b>[{_e(a.get("label"))}]</b></font> <b>{_e(a.get("judul"))}</b> — {_e(a.get("isi"))}', 6.6, markup=True)
                   for a in n.get("aksi", [])]
        for i, fl in enumerate((s_flows, a_flows)):
            box = Table([[fl]], colWidths=[half])
            box.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .4, hx(C["border"])), ("TOPPADDING", (0, 0), (-1, -1), 5),
                                     ("LEFTPADDING", (0, 0), (-1, -1), 7)]))
            box.wrap(half, max(y - 30, 40))
            box.drawOn(cv, 24 + i * (half + 12), max(y - box._height, 30))
        cv.setFont(regular, 5.6)
        cv.setFillColor(hx(C["muted"]))
        cv.drawString(24, 14, f"Gap = (AVG Harga {comp} − HTD PTPL+3%) / HTD PTPL+3% · Data Source: {ds.get('source_label', idm.SOURCE_LABEL)}, "
                              f"{pa['label']} · Confidential — Internal Use Only · {meta['company']}")
        cv.showPage()
    cv.save()
    return buf.getvalue()


# ==========================================================================
# PowerPoint (satu slide per tahap)
# ==========================================================================
def render_pptx(ds: dict, narr: dict, meta: dict) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt

    from .dashboard_render import logo_bytes, plain

    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(1000), Pt(720)
    rgb = lambda c: RGBColor.from_string(c)  # noqa: E731
    lb = logo_bytes()
    comp = meta["competitor"].title()
    pa, pb = ds["period_a"], ds["period_b"]
    F = "Arial"

    def text(s, x, y, w, h, runs, size=7, color=C["text"], bold=False, align=None):
        tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(1)
        for i, item in enumerate(runs if isinstance(runs, list) else [runs]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            parts = item if isinstance(item, list) else [(item, size, color, bold)]
            for t, sz, col, b in parts:
                r = p.add_run()
                r.text = plain(t)
                r.font.size, r.font.bold, r.font.name = Pt(sz), b, F
                r.font.color.rgb = rgb(col)
            if align:
                p.alignment = align
        return tb

    def rect(s, x, y, w, h, fill=None, line=None):
        from pptx.enum.shapes import MSO_SHAPE

        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
        r.shadow.inherit = False
        if fill:
            r.fill.solid()
            r.fill.fore_color.rgb = rgb(fill)
        else:
            r.fill.background()
        if line:
            r.line.color.rgb = rgb(line)
            r.line.width = Pt(.5)
        else:
            r.line.fill.background()
        return r

    def section(s, label, note, y):
        rect(s, 24, y, 2.5, 10, C["navy"])
        text(s, 31, y - 1, 400, 12, label.upper(), 7.5, C["navy"], True)
        if note:
            text(s, 500, y - 1, 476, 12, note, 6.3, C["muted"], align=PP_ALIGN.RIGHT)
        return y + 15

    for st in idm.STAGES:
        data = ds["stages"][st]
        s = prs.slides.add_slide(prs.slide_layouts[6])
        rect(s, 24, 16, 952, 2.5, C["navy"])
        if lb:
            s.shapes.add_picture(io.BytesIO(lb), Pt(24), Pt(24), height=Pt(30))
        text(s, 130, 22, 600, 20, "Price Competitiveness Analysis — B2B Segment", 14, C["navy"], True)
        text(s, 130, 42, 600, 12, f"{meta['company']} · {len(ds['segments'])} Segmen Customer Fokus · 3 Zona Wilayah · "
                                  f"{len({f['category'] for f in ds['focus']})} Kategori Produk", 7.5, C["muted"])
        text(s, 650, 22, 326, 34, [[(f"{pa['label']} (Most Recent)", 7, C["navy"], True)],
                                   [(f"Main stage: {idm.STAGE_LABEL[st]}", 6.3, C["muted"], False)],
                                   [(f"Gap positif = PTPL kompetitif | Gap negatif = {comp} lebih murah", 6.3, C["muted"], False)]],
             align=PP_ALIGN.RIGHT)
        y = 66
        if not data["has_data"]:
            text(s, 24, y + 10, 900, 20, f"Tidak ada data {idm.STAGE_LABEL[st]} untuk periode {pa['label']}.", 10, C["muted"])
            continue
        y = section(s, "A  Price Index Highlight", f"Current = {pa['label']}" + (f" | Trend = Δ vs {pb['label']}" if pb else ""), y)
        cards = kpi_cards(data, meta["competitor"])
        cw = (952 - 5 * 6) / 6
        for i, c in enumerate(cards):
            x = 24 + i * (cw + 6)
            rect(s, x, y, cw, 92, None, C["border"])
            rect(s, x, y, cw, 2.5, TONE_FG[c["tone"]])
            runs = [[(c["dim"], 5.8, C["muted"], True)]]
            if c["tag"]:
                runs.append([(c["tag"], 5.5, C["red"] if "NEG" in c["tag"] or "TIDAK" in c["tag"] else C["green"] if "TOP" in c["tag"] else C["amber"], True)])
            runs += [[(c["name"], 8, C["text"], True)],
                     [(pct(c["value"]), 15, TONE_FG[c["tone"]], False), ("  " + c["delta"], 6.5,
                                                                        C["green"] if c["dcls"] == "up" else C["red"] if c["dcls"] == "dn" else C["muted"], True)],
                     [(c["text"], 5.6, C["muted"], False)]]
            text(s, x + 4, y + 4, cw - 8, 86, runs)
        y += 100
        y = section(s, "B  Price Gap Matrix — Produk Fokus",
                    f"Current = Gap {pa['label'].split()[0]} vs HTD PTPL+3% | Sub = Δ vs {pb['label'].split()[0] if pb else '-'} | Angka = Max/Avg/Min Harga {comp}", y)
        cats = [c for c in idm.CATEGORY_ORDER if any(f["category"] == c for f in ds["focus"])] + \
            sorted({f["category"] for f in ds["focus"]} - set(idm.CATEGORY_ORDER))
        n_rows = 2 + len(cats) + len(ds["focus"])
        row_h = 9.5
        prod_h = 23
        total_h = 2 * 12 + len(cats) * row_h + len(ds["focus"]) * prod_h
        ncol = 4 + len(ds["segments"])
        tbl = s.shapes.add_table(n_rows, ncol, Pt(24), Pt(y), Pt(952), Pt(total_h)).table
        tbl.columns[0].width = Pt(92)
        for ci in range(1, ncol):
            tbl.columns[ci].width = Pt((952 - 92) / (ncol - 1))

        def fill(cell, runs, bg="FFFFFF", align=PP_ALIGN.CENTER):
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Pt(2)
            cell.margin_top = cell.margin_bottom = Pt(0)
            tf = cell.text_frame
            tf.word_wrap = True
            for i, parts in enumerate(runs):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = align
                for t, sz, col, b in parts:
                    r = p.add_run()
                    r.text = plain(t)
                    r.font.size, r.font.bold, r.font.name = Pt(sz), b, F
                    r.font.color.rgb = rgb(col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(bg)

        tbl.cell(0, 1).merge(tbl.cell(0, 3))
        tbl.cell(0, 4).merge(tbl.cell(0, ncol - 1))
        fill(tbl.cell(0, 0), [[("", 6, C["navy"], True)]], "FFFFFF")
        fill(tbl.cell(0, 1), [[("GAP HARGA PER ZONA WILAYAH VS PRODUK FOKUS", 6, "FFFFFF", True)]], C["navy"])
        fill(tbl.cell(0, 4), [[("GAP HARGA PER SEGMENT CUSTOMER VS PRODUK FOKUS", 6, C["amber"], True)]], C["amber_bg"])
        fill(tbl.cell(1, 0), [[("KATEGORI / PRODUK", 5.8, C["navy"], True)]], C["bg_alt"], PP_ALIGN.LEFT)
        for i, z in enumerate(idm.ZONES, start=1):
            fill(tbl.cell(1, i), [[(f"{z.upper()} {idm.ZONE_SUB[z]}", 5.6, C["navy"], True)]], C["bg_alt"])
        for i, ch in enumerate(ds["segments"], start=4):
            fill(tbl.cell(1, i), [[(ch.upper(), 5.6, C["navy"], True)]], C["bg_alt"])
        tbl.rows[0].height = tbl.rows[1].height = Pt(12)

        def cell_runs(cell):
            if not cell or cell.get("current") is None:
                return [[("—", 7, C["muted"], False)]], "FFFFFF"
            t = idm.tone(cell["current"])
            d, cls = delta(cell["current"], cell.get("prev"))
            runs = [[(pct(cell["current"]), 7, TONE_FG[t], True), ("  " + d if d else "", 5.3,
                                                                 C["green"] if cls == "up" else C["red"] if cls == "dn" else C["muted"], True)]]
            brands = ", ".join(b.title() for b in cell.get("brands") or [])
            runs.append([(f"{brands}  Max {idr(cell['max'])} Avg {idr(cell['avg'])} Min {idr(cell['min'])}", 4.4, C["muted"], False)])
            return runs, TONE_BG[t]

        r = 2
        for cat in cats:
            tbl.cell(r, 0).merge(tbl.cell(r, ncol - 1))
            fill(tbl.cell(r, 0), [[(cat.upper(), 6, C["navy"], True)]], C["bg_alt"], PP_ALIGN.LEFT)
            tbl.rows[r].height = Pt(row_h)
            r += 1
            for f in [f for f in ds["focus"] if f["category"] == cat]:
                p = f["name"]
                fill(tbl.cell(r, 0), [[(idm.short_name(p), 7, C["text"], False)]], "FFFFFF", PP_ALIGN.LEFT)
                for ci, cell in enumerate([data["table1"][p][z] for z in idm.ZONES] + [data["table2"][c][p] for c in ds["segments"]], start=1):
                    runs, bg = cell_runs(cell)
                    fill(tbl.cell(r, ci), runs, bg)
                tbl.rows[r].height = Pt(prod_h)
                r += 1
        y += total_h + 10
        y = section(s, "C  Executive Summary & Actionable Insights", "", y)
        n = narr.get(st) or {}
        half = (952 - 12) / 2
        box_h = max(720 - y - 26, 40)
        rect(s, 24, y, half, box_h, None, C["border"])
        rect(s, 24 + half + 12, y, half, box_h, None, C["border"])
        text(s, 30, y + 3, half - 12, box_h - 6, [[("EXECUTIVE SUMMARY", 7.5, C["navy"], True)]] +
             [[("• " + (r_.get("judul") or "") + " — ", 6.4, C["navy"], True), (r_.get("isi") or "", 6.4, C["text"], False)] for r_ in n.get("ringkasan", [])])
        text(s, 30 + half + 12, y + 3, half - 12, box_h - 6, [[("ACTIONABLE INSIGHTS", 7.5, C["navy"], True)]] +
             [[(f"[{a.get('label')}] ", 6.4, C["red"] if a.get("label") == "URGENT" else C["navy"] if a.get("label") == "LEVERAGE" else C["amber"], True),
               ((a.get("judul") or "") + " — ", 6.4, C["navy"], True), (a.get("isi") or "", 6.4, C["text"], False)] for a in n.get("aksi", [])])
        text(s, 24, 700, 952, 12, f"Gap = (AVG Harga {comp} − HTD PTPL+3%) / HTD PTPL+3% · Data Source: {ds.get('source_label', idm.SOURCE_LABEL)}, "
                                  f"{pa['label']} · Confidential — Internal Use Only · {meta['company']}", 5.6, C["muted"])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


RENDERERS = {"html": render_html, "pdf": render_pdf, "pptx": render_pptx}


# ==========================================================================
# Data ilustrasi (contoh tampilan & pengujian)
# ==========================================================================
def sample_rows(seed: int = 11) -> list[dict]:
    rnd = random.Random(seed)
    names = {"Meditran SX Plus 15W40": ("MEDITRAN SX PLUS 15W-40", 45500), "Meditran S": ("PERTAMINA MEDITRAN S 40", 36000),
             "Turalik 52": ("TURALIK 52", 35400), "Rored HDA 90": ("RORED HDA 90", 45300),
             "Masri RG 320": ("MASRI RG 320", 55100), "Medripal 412": ("MEDRIPAL 412", 41900),
             "Grease Pertamina SGX-NL 2": ("GREASE PERTAMINA SGX-NL 2", 69830)}
    rows = []
    for st in idm.STAGES:
        for per in ("A", "B"):
            for _name, (raw, htd) in names.items():
                for zone in idm.ZONES:
                    for ch in idm.CHANNELS:
                        for _ in range(rnd.randint(0, 2)):
                            factor = rnd.uniform(0.82, 1.45) * (0.97 if per == "B" else 1.0)
                            rows.append({"PERIOD": per, "STAGE": st, "ZONE": zone, "CHANNEL": ch, "PRODUCT": raw, "BRAND": "SHELL",
                                         "PRICE": round(htd * factor), "HTD_PLUS": htd})
    return rows

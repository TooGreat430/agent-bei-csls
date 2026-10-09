"""Render dashboard daya saing harga ke HTML (bertab), PDF (per halaman), dan PowerPoint (per slide).

Masukan: daftar halaman dari price_dashboard.build_pages() + meta (judul, perusahaan, tanggal, logo).
"""
from __future__ import annotations

import base64
import html as _html
import io
import math
import os
from typing import Any

from .report_render import fmt_id, nice_ticks

C = {"navy": "0F2347", "accent": "1E6FD9", "gold": "E0A526", "orange": "D9822B", "low": "1E7B4F", "mid": "B7791F",
     "high": "C0392B", "muted": "5B6470", "line": "D5DBE3", "light": "EEF2F7", "hero": "FFF6D6", "anom": "FFF7E0",
     "text": "1E2329", "white": "FFFFFF", "bar_b": "9AA9BF"}
ZONE_COLORS = {"Nasional": C["navy"], "Zona 1": C["accent"], "Zona 2": C["gold"], "Zona 3": C["low"]}
TONE = {"low": C["low"], "mid": C["mid"], "high": C["high"], "muted": C["muted"]}
LEVEL = {"positif": C["low"], "perhatian": C["mid"], "kritis": C["high"]}


def series_colors(chart: dict[str, Any]) -> list[str]:
    names = [s["nama"] for s in chart["seri"]]
    if all(n in ZONE_COLORS for n in names):
        return [ZONE_COLORS[n] for n in names]
    if chart.get("period_compare") and len(names) == 2:
        return [C["bar_b"], C["navy"]]
    if names and names[0].lower().startswith("margin ptpl"):
        return [C["low"], C["bar_b"]][:len(names)]
    return [C["navy"], C["accent"], C["gold"], C["low"]][:len(names)]


def plain(text: Any) -> str:
    """Teks aman untuk font PDF/PowerPoint: buang simbol yang tidak tersedia."""
    t = "" if text is None else str(text)
    for sym in ("✓ ", "⚠ ", "⚡ ", "★ ", "✓", "⚠", "⚡", "★"):
        t = t.replace(sym, "")
    return t.replace("▸", "›").strip()


def sec_text(title: Any) -> tuple[str, str]:
    if isinstance(title, (tuple, list)):
        return plain(title[0]), plain(title[1]) if len(title) > 1 else ""
    return plain(title), ""


def _e(text: Any) -> str:
    return _html.escape("" if text is None else str(text))


def logo_bytes() -> bytes | None:
    path = os.path.join(os.path.dirname(__file__), "assets", "logo_ptpl.png")
    if os.path.exists(path):
        with open(path, "rb") as fh:
            return fh.read()
    return None


# ==========================================================================
# HTML
# ==========================================================================
def svg_bars(chart: dict[str, Any], width: int = 420, height: int = 190) -> str:
    cats, series = chart["kategori"], chart["seri"]
    colors = ["#" + c for c in series_colors(chart)]
    vals = [v for s in series for v in s["nilai"] if v is not None]
    if not vals:
        return ""
    ticks = nice_ticks(min(0.0, min(vals)), max(0.0, max(vals)), 4)
    lo, hi = ticks[0], ticks[-1]
    left, right, top, bottom = 46, 6, 8, 44
    pw, ph = width - left - right, height - top - bottom

    def y(v):
        return top + ph - (v - lo) / (hi - lo) * ph

    out = [f'<svg viewBox="0 0 {width} {height}" xmlns="http://www.w3.org/2000/svg" style="width:100%;height:auto">']
    for t in ticks:
        out.append(f'<line x1="{left}" x2="{left + pw}" y1="{y(t):.1f}" y2="{y(t):.1f}" stroke="#{C["line"]}" '
                   f'stroke-width="{1.2 if t == 0 else .5}"/><text x="{left - 4}" y="{y(t) + 3:.1f}" text-anchor="end" '
                   f'font-size="8" fill="#{C["muted"]}">{fmt_id(t)}</text>')
    group = pw / len(cats)
    bw = group * 0.8 / len(series)
    for si, s in enumerate(series):
        for ci, v in enumerate(s["nilai"]):
            if v is None:
                continue
            x = left + ci * group + group * 0.1 + si * bw
            y0, y1 = sorted((y(0), y(v)))
            out.append(f'<rect x="{x:.1f}" y="{y0:.1f}" width="{bw * .9:.1f}" height="{max(y1 - y0, .8):.1f}" '
                       f'fill="{colors[si % len(colors)]}"><title>{_e(s["nama"])} · {_e(cats[ci])}: {fmt_id(v)}</title></rect>')
    for ci, c in enumerate(cats):
        words, lines, cur = c.split(), [], ""
        for w in words:
            if len(cur) + len(w) + 1 > 14 and cur:
                lines.append(cur)
                cur = w
            else:
                cur = (cur + " " + w).strip()
        lines.append(cur)
        cx = left + ci * group + group / 2
        for li, line in enumerate(lines[:3]):
            out.append(f'<text x="{cx:.1f}" y="{top + ph + 11 + li * 9}" text-anchor="middle" font-size="7.5" '
                       f'fill="#{C["text"]}">{_e(line)}</text>')
    lx = left
    for si, s in enumerate(series):
        out.append(f'<rect x="{lx}" y="{height - 9}" width="8" height="8" fill="{colors[si % len(colors)]}"/>'
                   f'<text x="{lx + 11}" y="{height - 2}" font-size="8" fill="#{C["text"]}">{_e(s["nama"])}</text>')
        lx += 20 + 4.6 * len(s["nama"])
    out.append("</svg>")
    return "".join(out)


def _cell_html(value: str, tone: str | None) -> str:
    if tone and tone.startswith("status:"):
        t = tone.split(":")[1]
        return f'<td class="st"><span class="dot" style="background:#{TONE.get(t, C["muted"])}"></span>{_e(value)}</td>'
    style = f' style="color:#{TONE[tone]}"' if tone in TONE else ""
    return f"<td{style}>{_e(value)}</td>"


def _sec(title: tuple | str, red: bool = False) -> str:
    main, accent = (title if isinstance(title, tuple) else (title, ""))
    acc = f' <span class="acc">· {_e(accent)}</span>' if accent else ""
    return f'<div class="sec{" red" if red else ""}">{_e(main)}{acc}</div>'


def _table_html(table: dict[str, Any]) -> str:
    cols = table["columns"]
    if table.get("groups"):
        top = "".join(f'<th colspan="{span}" class="gh">{_e(label)}</th>' for label, span in table["groups"])
        head = f"<tr>{top}</tr><tr class='sub'>" + "".join(f"<th>{_e(c)}</th>" for c in cols) + "</tr>"
    else:
        head = "<tr>" + "".join(f"<th>{_e(c)}</th>" for c in cols) + "</tr>"
    body = ""
    for r in table["rows"]:
        if r["kind"] == "group":
            body += f'<tr class="grp"><td colspan="{len(cols)}">{_e(r["cells"][0])}</td></tr>'
            continue
        body += f'<tr class="{r["kind"]}">' + "".join(_cell_html(v, t) for v, t in zip(r["cells"], r["tones"])) + "</tr>"
    note = f'<div class="tnote">{_e(table.get("note"))}</div>' if table.get("note") else ""
    return f'<div class="twrap"><table>{head}{body}</table></div>{note}'


def _insights_html(items: list[dict]) -> str:
    return "".join(f'<div class="ins" style="border-color:#{LEVEL.get(i.get("tingkat"), C["accent"])}">'
                   f'<small>{_e(i.get("kategori"))}</small><b>{_e(i.get("judul"))}</b>{_e(i.get("uraian"))}</div>'
                   for i in items)


def _watch_html(items: list[dict]) -> str:
    out = ""
    for w in items:
        stats = "".join(f'<div class="ws"><small>{_e(lbl)}</small><b style="color:#{TONE.get(t, C["text"])}">{_e(v)}</b></div>'
                        for lbl, v, t in w["stats"])
        months = "".join(
            f'<div class="wmo {("g" if t == "low" else "r" if t == "high" else "")}"><small>{_e(m)}</small>'
            f'<b>{_e(v)}</b></div>' for m, v, t in zip(w["months"], w["values"], w["tones"]))
        brand = f'<span class="brand">{_e(w["brand"])}</span>' if w["brand"] else ""
        out += (f'<div class="watch"><div class="wt"><span>{_e(w["product"])}</span>{brand}</div>'
                f'<div class="wm">{_e(w.get("subtitle") or (w["segment"] + " · " + w["metric"] + " = PTPL − produk ini (IDR/L)"))}</div>'
                f'<div class="wstats">{stats}</div><div class="wmonths">{months}</div>'
                f'<div class="wcap">{_e(w["caption"])}</div></div>')
    return out


def _chip_html(text: str, tone: str) -> str:
    return f'<span class="chip c-{tone}">{_e(text)}</span>'


def _page_html(page: dict[str, Any], meta: dict[str, Any], logo: str) -> str:
    right = _e(meta["period_text"]) + (f'<br>{_e(page["zone_label"])}' if page.get("zone_label") else "")
    head = (f'<div class="ph"><img src="{logo}" alt="Pertamina Lubricants" class="logo"/>'
            f'<div class="pt"><h2>{_e(page["title"])}</h2><span class="tag">{_e(page["tag"])}</span>'
            f'<p>{_e(page["subtitle"])}</p></div><div class="pm">{right}<br>'
            f'<span class="conf">CONFIDENTIAL · C-SUITE ANALYTICS</span></div></div>')
    sec = page.get("sec", {})
    if page["kind"] == "exec":
        cards = ""
        for i, z in enumerate(page["zone_cards"]):
            color = list(ZONE_COLORS.values())[i % 4]
            rows = "".join(f'<div class="zm"><span>{_e(m["label"])}</span><b style="color:#{TONE.get(m["tone"], C["text"])}">'
                           f'{_e(m["value"])}</b></div>' for m in z["metrics"])
            note = f'<div class="kn"><b>{_e(z["note_label"])}</b> {_e(z["note"])}</div>' if z["note"] else ""
            cards += (f'<div class="kpi zc" style="border-top-color:#{color}"><div class="zh"><span>{_e(z["title"])}</span>'
                      f'<span class="nick">{_e(z["subtitle"])}</span></div>{rows}{note}</div>')
        bullets = ("<ul class='sum'>" + "".join(f"<li>{_e(b)}</li>" for b in page["bullets"]) + "</ul>") if page["bullets"] else ""
        charts = "".join(f'<div class="ch"><div class="cht">{_e(c["judul"])} ({_e(c["satuan"])})</div>{svg_bars(c)}</div>'
                         for c in page["charts"])
        body = (f'<div class="kpis">{cards}</div>{bullets}<div class="grid2"><div>{_sec(sec["table"])}{_table_html(page["table"])}</div>'
                f'<div>{_sec(sec["insights"])}{_insights_html(page["insights"])}<div class="charts2">{charts}</div></div></div>')
    else:
        kpis = "".join(
            f'<div class="kpi" style="border-top-color:#{TONE.get(k["chip_tone"], C["accent"])}"><div class="kl">{_e(k["label"])}</div>'
            f'<div class="kv" style="color:#{TONE.get(k["tone"], C["text"])}">{_e(k["value"])}</div>'
            f'<div class="kd">{_e(k["detail"])}</div><div class="kn">{_e(k["extra"])}</div>{_chip_html(k["chip"], k["chip_tone"])}</div>'
            for k in page["kpis"])
        charts = "".join(f'<div class="ch"><div class="cht">{_e(c["judul"])}</div>{svg_bars(c, 380, 170)}</div>'
                         for c in page["charts"])
        anomaly = (f"<div class='anom'><b>⚠ {_e(page['anomaly_title'])}</b><ul>" + "".join(f"<li>{_e(a)}</li>" for a in page["anomaly"])
                   + "</ul></div>") if page["anomaly"] else ""
        watch = (_sec(sec["watch"], red=True) + _watch_html(page["watch"])) if page["watch"] else ""
        body = (f'<div class="kpis">{kpis}</div><div class="grid3"><div>{_sec(sec["table"])}{_table_html(page["table"])}</div>'
                f'<div>{_sec(sec["charts"])}{charts}{anomaly}</div>'
                f'<div>{_sec(sec["insights"])}{_insights_html(page["insights"])}{watch}</div></div>'
                + (f'<div class="fnote">{_e(page["footnote"])}</div>' if page.get("footnote") else ""))
    return f'<section class="page">{head}{body}<div class="pf">{_e(page["footer"])}</div></section>'


HTML_CSS = f"""
*{{box-sizing:border-box}}body{{margin:0;background:#E9EDF2;font-family:'Liberation Sans',Arial,sans-serif;color:#{C['text']};font-size:12px}}
.top{{position:sticky;top:0;z-index:5;background:#0B1B33;color:#fff;display:flex;align-items:center;justify-content:space-between;padding:10px 22px;gap:16px}}
.top h1{{font-size:13px;margin:0;letter-spacing:.3px;flex:1;min-width:260px}}.tabs{{display:flex;gap:5px;flex-wrap:wrap;justify-content:flex-end}}
.tabs button{{background:#1C2E4E;color:#fff;border:1px solid #2F4569;border-radius:4px;padding:6px 9px;font-size:10.5px;font-weight:700;cursor:pointer;white-space:nowrap}}
.tabs button.on{{background:#{C['gold']};color:#0B1B33;border-color:#{C['gold']}}}
.tab{{display:none;padding:18px}}.tab.on{{display:block}}
.page{{background:#fff;max-width:1400px;margin:0 auto 18px;border-radius:6px;box-shadow:0 1px 4px rgba(0,0,0,.08);padding:16px 18px}}
.ph{{display:flex;gap:14px;align-items:center;border-bottom:2px solid #{C['navy']};padding-bottom:10px}}.logo{{height:38px}}
.pt{{flex:1}}.pt h2{{margin:0;font-size:17px;color:#{C['navy']}}}.pt p{{margin:3px 0 0;color:#{C['orange']};font-size:11.5px;font-weight:700}}
.tag{{display:inline-block;background:#{C['navy']};color:#fff;font-size:9px;font-weight:700;border-radius:3px;padding:2px 6px;margin-top:3px}}
.pm{{text-align:right;font-size:10px;color:#{C['muted']};line-height:1.6}}.conf{{background:#{C['navy']};color:#fff;padding:2px 6px;border-radius:3px;font-weight:700}}
.kpis{{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin:12px 0}}
.kpi{{border:1px solid #{C['line']};border-top:3px solid #{C['accent']};border-radius:4px;padding:8px 10px}}
.kl{{font-size:10px;font-weight:700;color:#{C['muted']};text-transform:uppercase}}.kl small{{text-transform:none;font-weight:400}}
.kv{{font-size:20px;font-weight:700;margin:2px 0}}.kd{{font-family:'Liberation Mono',monospace;font-size:10px;color:#{C['accent']}}}
.kn{{font-size:10px;color:#{C['muted']};margin:3px 0 5px}}
.zm{{display:flex;justify-content:space-between;font-size:11px;padding:2px 0;border-bottom:1px dashed #{C['line']}}}
.badge{{display:inline-block;color:#fff;font-size:9px;font-weight:700;border-radius:3px;padding:2px 6px}}
.grid3{{display:grid;grid-template-columns:45% 28% 27%;gap:12px}}.grid2{{display:grid;grid-template-columns:58% 42%;gap:14px}}
.sec{{font-size:11px;font-weight:700;color:#{C['navy']};text-transform:uppercase;border-bottom:1px solid #{C['line']};padding-bottom:3px;margin:4px 0 6px}}
.sec.red{{color:#{C['high']};margin-top:10px}}
.twrap{{overflow-x:auto}}table{{border-collapse:collapse;width:100%;font-size:10.5px}}
th{{background:#{C['navy']};color:#fff;padding:4px 5px;text-align:right;font-size:9.5px}}th:first-child{{text-align:left}}
td{{padding:3px 5px;border-bottom:1px solid #EDF0F4;text-align:right;white-space:nowrap}}td:first-child{{text-align:left;white-space:normal}}
td.st{{text-align:center}}tr.grp td{{background:#{C['light']};font-weight:700;color:#{C['navy']};text-align:left}}
tr.hero td{{background:#{C['hero']}}}tr.hero td:first-child{{font-weight:700}}
.tnote{{font-size:9.5px;color:#{C['muted']};margin-top:4px}}
.ch{{margin-bottom:8px}}.cht{{font-size:10.5px;font-weight:700;color:#{C['navy']};margin-bottom:2px}}
.anom{{background:#{C['anom']};border:1px solid #{C['gold']};border-radius:4px;padding:8px 10px;font-size:10.5px}}.anom ul{{margin:4px 0 0;padding-left:16px}}
.ins{{border-left:4px solid;background:#F7F9FB;padding:7px 9px;margin-bottom:7px;font-size:10.5px;line-height:1.4}}
.ins small{{display:block;font-size:9px;font-weight:700;color:#{C['muted']}}}.ins b{{display:block;color:#{C['navy']};margin:1px 0}}
.watch{{border:1px solid #{C['line']};border-radius:4px;padding:6px 8px;margin-bottom:6px}}.wt{{font-weight:700;font-size:10.5px;display:flex;justify-content:space-between;gap:6px}}
.wm{{font-size:9px;color:#{C['muted']}}}.wtab td{{text-align:center;font-size:10px;border:0;padding:3px 2px}}.wtab small{{display:block;font-size:8.5px;color:#{C['muted']}}}
.sum{{margin:0 0 10px;padding-left:18px;font-size:11.5px;line-height:1.5}}.charts2{{display:grid;grid-template-columns:1fr 1fr;gap:8px;margin-top:6px}}
.pf{{font-size:9.5px;color:#{C['muted']};border-top:1px solid #{C['line']};margin-top:10px;padding-top:6px}}
.acc{{color:#{C['orange']}}}.sec .acc{{font-weight:700}}
th.gh{{text-align:center;border-left:1px solid #2F4569}}th.gh:first-child{{text-align:left}}tr.sub th{{background:#EDF0F4;color:#{C['navy']};font-size:9px}}
tr.grp td{{background:#{C['navy']} !important;color:#fff !important;font-size:10px;letter-spacing:.3px}}
tr.hero td{{background:#{C['hero']}}}tr.hero td:first-child{{font-style:italic;color:#9A6A00;font-weight:400;padding-left:12px}}
tr.viscosity td:first-child{{font-weight:700}}
td.st{{text-align:left;font-weight:700;font-size:9.5px}}.dot{{display:inline-block;width:7px;height:7px;border-radius:50%;margin-right:4px}}
.chip{{display:inline-block;font-size:9px;font-weight:700;border-radius:3px;padding:2px 7px;border:1px solid}}
.c-low{{background:#E8F5EC;color:#{C['low']};border-color:#9FD3B2}}.c-mid{{background:#FFF4D6;color:#8A5A00;border-color:#EBC86B}}
.c-high{{background:#FDECEA;color:#{C['high']};border-color:#F1A9A1}}.c-muted{{background:#F2F4F7;color:#{C['muted']};border-color:#D5DBE3}}
.kd{{color:#B7791F !important;font-weight:700}}
.zh{{display:flex;justify-content:space-between;align-items:center;font-weight:700;color:#{C['navy']};font-size:11.5px;margin-bottom:6px}}
.nick{{background:#{C['navy']};color:#fff;font-size:9px;border-radius:3px;padding:2px 6px}}
.zc .kn{{border-top:1px dashed #{C['line']};padding-top:5px}}.zc .kn b{{color:#{C['navy']}}}
.watch{{background:#F4FBF6}}.wt{{align-items:center}}.brand{{background:#{C['high']};color:#fff;font-size:8.5px;border-radius:3px;padding:1px 6px}}
.wstats{{display:grid;grid-template-columns:repeat(3,1fr);gap:4px;margin:5px 0}}
.ws{{background:#EEF2F7;border-radius:3px;text-align:center;padding:3px}}.ws small{{display:block;font-size:8px;color:#{C['muted']}}}.ws b{{font-size:10.5px}}
.wmonths{{display:grid;grid-template-columns:repeat(6,1fr);gap:3px}}
.wmo{{border:1px solid #D5DBE3;border-radius:3px;text-align:center;padding:2px;font-size:9.5px}}.wmo small{{display:block;font-size:7.5px;color:#{C['muted']}}}
.wmo.g{{background:#E8F5EC;border-color:#9FD3B2;color:#{C['low']}}}.wmo.r{{background:#FDECEA;border-color:#F1A9A1;color:#{C['high']}}}
.wcap{{font-size:9px;font-style:italic;color:#{C['muted']};margin-top:3px}}
.fnote{{font-size:9.5px;color:#{C['muted']};margin-top:8px}}
"""


def render_html(pages: list[dict[str, Any]], meta: dict[str, Any]) -> bytes:
    lb = logo_bytes()
    logo = f"data:image/png;base64,{base64.b64encode(lb).decode()}" if lb else ""
    tabs, order = {}, []
    for p in pages:
        if p["tab"] not in tabs:
            tabs[p["tab"]] = []
            order.append(p["tab"])
        tabs[p["tab"]].append(p)
    buttons = "".join(f'<button data-t="{i}" class="{"on" if i == 0 else ""}">{_e(t)}</button>' for i, t in enumerate(order))
    sections = "".join(f'<div class="tab{" on" if i == 0 else ""}" id="t{i}">' + "".join(_page_html(p, meta, logo) for p in tabs[t])
                       + "</div>" for i, t in enumerate(order))
    script = ("document.querySelectorAll('.tabs button').forEach(b=>b.onclick=()=>{document.querySelectorAll('.tabs button')"
              ".forEach(x=>x.classList.remove('on'));document.querySelectorAll('.tab').forEach(x=>x.classList.remove('on'));"
              "b.classList.add('on');document.getElementById('t'+b.dataset.t).classList.add('on');window.scrollTo(0,0)})")
    doc = (f"<!DOCTYPE html><html lang='id'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
           f"<title>{_e(meta['report_title'])}</title><style>{HTML_CSS}</style></head><body>"
           f"<div class='top'><h1>{_e(pages[0].get('topbar') or meta['report_title'].upper())}</h1><div class='tabs'>{buttons}</div></div>"
           f"{sections}<script>{script}</script></body></html>")
    return doc.encode("utf-8")


# ==========================================================================
# PDF (reportlab) — satu halaman per bagian
# ==========================================================================
def render_pdf(pages: list[dict[str, Any]], meta: dict[str, Any]) -> bytes:
    from reportlab.graphics import renderPDF
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph, Table, TableStyle

    from .report_render import _register_fonts

    regular, bold, _italic = _register_fonts()
    W, H = 960, 540
    buf = io.BytesIO()
    cv = canvas.Canvas(buf, pagesize=(W, H))
    cv.setTitle(meta["report_title"])
    hx = lambda c: colors.HexColor("#" + c)  # noqa: E731
    lb = logo_bytes()
    logo = ImageReader(io.BytesIO(lb)) if lb else None

    def para(text, size=7, color=C["text"], font=None, leading=None, align=0, markup=False, italic=False):
        st = ParagraphStyle("p", fontName=font or regular, fontSize=size, leading=leading or size * 1.3,
                            textColor=hx(color), alignment=align)
        body = plain(text) if markup else _e(plain(text))
        return Paragraph(f"<i>{body}</i>" if italic else body, st)

    def put(flow, x, y_top, w, max_h=None):
        _, h = flow.wrap(w, max_h or H)
        flow.drawOn(cv, x, y_top - h)
        return h

    def header(page):
        if logo:
            cv.drawImage(logo, 22, H - 52, width=96, height=31, mask="auto")
        th = put(para(page["title"], 11.5, C["navy"], bold, leading=13), 128, H - 14, 560)
        put(para(page["subtitle"], 7.2, C["orange"], bold), 128, H - 16 - th, 560)
        cv.setFillColor(hx(C["navy"]))
        cv.roundRect(W - 182, H - 50, 160, 13, 2, fill=1, stroke=0)
        cv.setFillColor(colors.white)
        cv.setFont(bold, 6.5)
        cv.drawCentredString(W - 102, H - 46, "CONFIDENTIAL · C-SUITE ANALYTICS")
        cv.setFillColor(hx(C["muted"]))
        cv.setFont(regular, 6.5)
        tag = plain(page["tag"])
        tw = cv.stringWidth(tag, bold, 6) + 10
        cv.setFillColor(hx(C["navy"]))
        cv.roundRect(W - 22 - tw, H - 20, tw, 10, 1.5, fill=1, stroke=0)
        cv.setFillColor(colors.white)
        cv.setFont(bold, 6)
        cv.drawRightString(W - 27, H - 17, tag)
        cv.setFillColor(hx(C["muted"]))
        cv.setFont(regular, 6)
        cv.drawRightString(W - 22, H - 28, meta["period_text"])
        if page.get("zone_label"):
            cv.setFont(bold, 6.2)
            cv.setFillColor(hx(C["navy"]))
            cv.drawRightString(W - 22, H - 36, plain(page["zone_label"]))
        cv.setStrokeColor(hx(C["navy"]))
        cv.setLineWidth(1.2)
        cv.line(22, H - 58, W - 22, H - 58)

    def footer(page, n):
        cv.setFillColor(hx(C["muted"]))
        cv.setFont(regular, 5.8)
        cv.drawString(22, 12, page["footer"][:230])
        cv.drawRightString(W - 22, 12, str(n))

    def section(title, x, y, w, color=C["navy"]):
        main, accent = sec_text(title)
        full = main.upper() + ("  · " + accent.upper() if accent else "")
        size = 7.0
        while size > 5 and cv.stringWidth(full, bold, size) > w:
            size -= 0.3
        cv.setFillColor(hx(color))
        cv.setFont(bold, size)
        cv.drawString(x, y - 7, main.upper())
        if accent:
            cv.setFillColor(hx(C["orange"]))
            cv.drawString(x + cv.stringWidth(main.upper() + "  ", bold, size), y - 7, "· " + accent.upper())
        cv.setStrokeColor(hx(C["line"]))
        cv.setLineWidth(.5)
        cv.line(x, y - 10, x + w, y - 10)
        return y - 14

    def table(tbl, x, y_top, w, max_h):
        cols = tbl["columns"]
        first = w * (0.30 if len(cols) <= 9 else 0.22)
        widths = [first] + [(w - first) / (len(cols) - 1)] * (len(cols) - 1)
        rows = tbl["rows"]
        for size in (6.3, 5.9, 5.5, 5.1):
            data, cmds = _pdf_table_data(tbl, rows, size, para, regular, bold)
            t = Table(data, colWidths=widths)
            t.setStyle(TableStyle(cmds))
            _, h = t.wrap(w, max_h)
            if h <= max_h:
                break
        else:
            keep = max(4, int(len(rows) * max_h / h))
            data, cmds = _pdf_table_data(tbl, rows[:keep], 5.1, para, regular, bold)
            t = Table(data, colWidths=widths)
            t.setStyle(TableStyle(cmds))
            _, h = t.wrap(w, max_h)
        t.drawOn(cv, x, y_top - h)
        y = y_top - h - 3
        if tbl.get("note"):
            y -= put(para(tbl["note"], 5.6, C["muted"]), x, y, w)
        return y

    def chart(ch, x, y_top, w, h):
        cols = series_colors(ch)
        put(para(f"{ch['judul']} ({ch['satuan']})", 6.5, C["navy"], bold), x, y_top, w)
        vals = [v for s in ch["seri"] for v in s["nilai"] if v is not None]
        if not vals:
            return
        d = Drawing(w, h - 10)
        bc = VerticalBarChart()
        bc.x, bc.y, bc.width, bc.height = 34, 26, w - 40, h - 48
        bc.data = [tuple(v if v is not None else 0 for v in s["nilai"]) for s in ch["seri"]]
        ticks = nice_ticks(min(0.0, min(vals)), max(0.0, max(vals)), 4)
        bc.valueAxis.valueMin, bc.valueAxis.valueMax = ticks[0], ticks[-1]
        bc.valueAxis.valueStep = ticks[1] - ticks[0]
        bc.valueAxis.labelTextFormat = lambda v: fmt_id(v)
        bc.valueAxis.labels.fontName, bc.valueAxis.labels.fontSize = regular, 5.5
        bc.valueAxis.visibleGrid, bc.valueAxis.gridStrokeColor, bc.valueAxis.gridStrokeWidth = 1, hx(C["line"]), .3
        bc.categoryAxis.categoryNames = [_wrap(c, 11 if len(ch["kategori"]) > 3 else 14, 2) for c in ch["kategori"]]
        bc.categoryAxis.labels.fontName, bc.categoryAxis.labels.fontSize = regular, 5
        bc.categoryAxis.labels.boxAnchor, bc.categoryAxis.labelAxisMode = "n", "low"
        bc.groupSpacing, bc.barSpacing = 6, 0.5
        for i in range(len(ch["seri"])):
            bc.bars[i].fillColor, bc.bars[i].strokeColor = hx(cols[i % len(cols)]), None
        d.add(bc)
        lx = 34
        for i, s in enumerate(ch["seri"]):
            d.add(Rect(lx, 2, 6, 6, fillColor=hx(cols[i % len(cols)]), strokeColor=None))
            d.add(String(lx + 9, 2.5, s["nama"], fontName=regular, fontSize=5.5))
            lx += 16 + 2.9 * len(s["nama"])
        renderPDF.draw(d, cv, x, y_top - h)

    def cards(items, x, y_top, w, y_min):
        y = y_top
        for it in items:
            flows = [para(it.get("kategori"), 5.5, C["muted"], bold), para(it.get("judul"), 7, C["navy"], bold),
                     para(it.get("uraian"), 6.3, C["text"])]
            t = Table([[flows]], colWidths=[w])
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), hx("F7F9FB")),
                                   ("LINEBEFORE", (0, 0), (0, -1), 3, hx(LEVEL.get(it.get("tingkat"), C["accent"]))),
                                   ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 6)]))
            _, h = t.wrap(w, H)
            if y - h < y_min:
                break
            t.drawOn(cv, x, y - h)
            y -= h + 4
        return y

    def watch(items, title, x, y_top, w, y_min):
        if not items:
            return y_top
        y = section(title, x, y_top, w, C["high"])
        for it in items:
            head = Table([[para(it["product"], 6.3, C["navy"], bold), para(it["brand"], 5.5, C["white"], bold, align=1)]],
                         colWidths=[w - 40, 34])
            head.setStyle(TableStyle([("BACKGROUND", (1, 0), (1, 0), hx(C["high"])), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                      ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
            sub = para(it.get("subtitle") or f"{it['segment']} · {it['metric']} = PTPL − produk ini (IDR/L)", 5.2, C["muted"])
            st = Table([[para(lbl, 4.8, C["muted"], align=1) for lbl, _, _ in it["stats"]],
                        [para(v, 6, TONE.get(t, C["text"]), bold, align=1) for _, v, t in it["stats"]]],
                       colWidths=[(w - 6) / max(len(it["stats"]), 1)] * len(it["stats"]))
            st.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), hx(C["light"])), ("TOPPADDING", (0, 0), (-1, -1), 0),
                                    ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
            mt = Table([[para(m, 4.8, C["muted"], align=1) for m in it["months"]],
                        [para(v, 5.6, TONE.get(t, C["text"]), bold, align=1) for v, t in zip(it["values"], it["tones"])]],
                       colWidths=[(w - 6) / max(len(it["months"]), 1)] * len(it["months"]))
            cmds = [("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                    ("GRID", (0, 0), (-1, -1), .3, colors.white)]
            for ci, t in enumerate(it["tones"]):
                if t in ("low", "high"):
                    cmds.append(("BACKGROUND", (ci, 0), (ci, 1), hx("E8F5EC" if t == "low" else "FDECEA")))
            mt.setStyle(TableStyle(cmds))
            box = Table([[[head, sub, st, mt, para(it["caption"], 5, C["muted"])]]], colWidths=[w])
            box.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .5, hx(C["line"])), ("BACKGROUND", (0, 0), (-1, -1), hx("F4FBF6")),
                                     ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            _, h = box.wrap(w, H)
            if y - h < y_min:
                break
            box.drawOn(cv, x, y - h)
            y -= h + 3
        return y

    for n, page in enumerate(pages, start=1):
        header(page)
        top = H - 66
        if page["kind"] == "exec":
            cw = (W - 44 - 24) / 4
            card_h = 0.0
            for i, z in enumerate(page["zone_cards"]):
                x = 22 + i * (cw + 8)
                flows = [Table([[para(z["title"], 6.8, C["navy"], bold), para(z["subtitle"], 5.4, C["white"], bold, align=1)]],
                               colWidths=[cw - 84, 70], style=TableStyle([("BACKGROUND", (1, 0), (1, 0), hx(C["navy"])),
                                                                          ("TOPPADDING", (0, 0), (-1, -1), 0),
                                                                          ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                                                                          ("LEFTPADDING", (0, 0), (0, 0), 0)]))]
                rows = [[para(m["label"], 6, C["text"]), para(m["value"], 6.5, TONE.get(m["tone"], C["text"]), bold, align=2)]
                        for m in z["metrics"]]
                mt = Table(rows, colWidths=[cw * .62 - 8, cw * .38 - 8])
                mt.setStyle(TableStyle([("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                                        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
                note = [para(f"<b>{_e(z['note_label'])}</b> {_e(z['note'])}", 5.6, C["muted"], markup=True)] if z["note"] else []
                cell = Table([[flows + [mt] + note]], colWidths=[cw])
                cell.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .5, hx(C["line"])),
                                          ("LINEABOVE", (0, 0), (-1, 0), 2.5, hx(list(ZONE_COLORS.values())[i])),
                                          ("TOPPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 6)]))
                cell.wrap(cw, 120)
                cell.drawOn(cv, x, top - cell._height)
                card_h = max(card_h, cell._height)
            y = top - card_h - 10
            if page["bullets"]:
                for b in page["bullets"][:4]:
                    y -= put(para("• " + b, 6.6), 22, y, W - 44) + 1
                y -= 4
            lw = (W - 44) * .58
            ly = section(page["sec"]["table"], 22, y, lw)
            table(page["table"], 22, ly, lw, ly - 26)
            rx, rw = 22 + lw + 14, (W - 44) - lw - 14
            ry = section(page["sec"]["insights"], rx, y, rw)
            ry = cards(page["insights"], rx, ry, rw, 150)
            cw2 = (rw - 8) / 2
            ch_top = ry - 4
            for i, ch in enumerate(page["charts"][:2]):
                chart(ch, rx + i * (cw2 + 8), ch_top, cw2, min(ch_top - 24, 170))
        else:
            kw = (W - 44 - 24) / 4
            for i, k in enumerate(page["kpis"]):
                x = 22 + i * (kw + 8)
                chip = Table([[para(k["chip"], 5.5, TONE.get(k["chip_tone"], C["muted"]), bold, align=1)]], colWidths=[66])
                chip.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .5, hx(TONE.get(k["chip_tone"], C["muted"]))),
                                          ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
                chip.hAlign = "LEFT"
                flows = [para(k["label"], 5.8, C["muted"], bold), para(k["value"], 12, TONE.get(k["tone"], C["text"]), bold),
                         para(k["detail"], 5.6, C["mid"], bold), para(k["extra"], 5.4, C["muted"]), chip]
                cell = Table([[flows]], colWidths=[kw])
                cell.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .5, hx(C["line"])),
                                          ("LINEABOVE", (0, 0), (-1, 0), 2.5, hx(TONE.get(k["chip_tone"], C["accent"]))),
                                          ("TOPPADDING", (0, 0), (-1, -1), 3), ("LEFTPADDING", (0, 0), (-1, -1), 6)]))
                cell.wrap(kw, 80)
                cell.drawOn(cv, x, top - cell._height)
            y = top - 76
            w1, w2 = (W - 44) * .45, (W - 44) * .28
            w3 = (W - 44) - w1 - w2 - 20
            x1, x2, x3 = 22, 22 + w1 + 10, 22 + w1 + w2 + 20
            ty = section(page["sec"]["table"], x1, y, w1)
            table(page["table"], x1, ty, w1, ty - 30)
            cy = section(page["sec"]["charts"], x2, y, w2)
            n_ch = max(len(page["charts"]), 1)
            anom_h = 16 + 10 * len(page["anomaly"]) if page["anomaly"] else 0
            ch_h = min(118, (cy - 26 - anom_h) / n_ch)
            for ch in page["charts"]:
                chart(ch, x2, cy, w2, ch_h)
                cy -= ch_h + 2
            if page["anomaly"]:
                flows = [para(page["anomaly_title"], 6.3, C["mid"], bold)] + [para("• " + a, 6, C["text"]) for a in page["anomaly"]]
                box = Table([[flows]], colWidths=[w2])
                box.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), hx(C["anom"])), ("BOX", (0, 0), (-1, -1), .6, hx(C["gold"])),
                                         ("TOPPADDING", (0, 0), (-1, -1), 4)]))
                box.wrap(w2, H)
                box.drawOn(cv, x2, max(cy - box._height, 34))
            iy = section(page["sec"]["insights"], x3, y, w3)
            iy = cards(page["insights"], x3, iy, w3, 230)
            watch(page["watch"], page["sec"]["watch"], x3, iy - 2, w3, 30)
            if page.get("footnote"):
                cv.setFont(regular, 5.6)
                cv.setFillColor(hx(C["muted"]))
                cv.drawString(22, 22, plain(page["footnote"])[:220])
        footer(page, n)
        cv.showPage()
    cv.save()
    return buf.getvalue()


def _wrap(text: str, width: int, max_lines: int) -> str:
    words, lines, cur = str(text).split(), [], ""
    for w in words:
        if cur and len(cur) + len(w) + 1 > width:
            lines.append(cur)
            cur = w
        else:
            cur = (cur + " " + w).strip()
    lines.append(cur)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1][: width - 1] + "…"
    return "\n".join(lines)


def _pdf_table_data(tbl, rows, size, para, regular, bold):
    from reportlab.lib import colors

    hx = lambda c: colors.HexColor("#" + c)  # noqa: E731
    ncol = len(tbl["columns"])
    data, cmds = [], [("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                      ("TOPPADDING", (0, 0), (-1, -1), .8), ("BOTTOMPADDING", (0, 0), (-1, -1), .8),
                      ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                      ("LINEBELOW", (0, 0), (-1, -1), .25, hx("E6EAF0"))]
    if tbl.get("groups"):
        top, col = [], 0
        for label, span in tbl["groups"]:
            top += [para(label, size - .4, C["white"], bold, align=0 if col == 0 else 1)] + [""] * (span - 1)
            if span > 1:
                cmds.append(("SPAN", (col, 0), (col + span - 1, 0)))
            col += span
        data.append(top)
        data.append([para(c, size - .6, C["navy"], bold, align=0 if i == 0 else 2) for i, c in enumerate(tbl["columns"])])
        cmds += [("BACKGROUND", (0, 0), (-1, 0), hx(C["navy"])), ("BACKGROUND", (0, 1), (-1, 1), hx("EDF0F4"))]
    else:
        data.append([para(c, size - .3, C["white"], bold, align=0 if i == 0 else 2) for i, c in enumerate(tbl["columns"])])
        cmds.append(("BACKGROUND", (0, 0), (-1, 0), hx(C["navy"])))
    offset = len(data)
    for ri, r in enumerate(rows, start=offset):
        if r["kind"] == "group":
            data.append([para(r["cells"][0], size - .2, C["white"], bold)] + [""] * (ncol - 1))
            cmds += [("SPAN", (0, ri), (-1, ri)), ("BACKGROUND", (0, ri), (-1, ri), hx(C["navy"]))]
            continue
        line = []
        for ci, (v, t) in enumerate(zip(r["cells"], r["tones"])):
            if t and t.startswith("status:"):
                color = TONE.get(t.split(":")[1], C["muted"])
                line.append(para(f'<font color="#{color}">●</font> {_e(v)}', size - .5, C["text"], bold, markup=True))
            elif ci == 0:
                line.append(para(v, size, "9A6A00" if r["kind"] == "hero" else C["text"],
                                 None if r["kind"] == "hero" else bold, italic=r["kind"] == "hero"))
            else:
                line.append(para(v, size, TONE.get(t, C["text"]), bold if t in TONE else None, align=2))
        data.append(line)
        if r["kind"] == "hero":
            cmds.append(("BACKGROUND", (0, ri), (-1, ri), hx(C["hero"])))
    return data, cmds


# ==========================================================================
# PowerPoint — satu slide per bagian
# ==========================================================================
def render_pptx(pages: list[dict[str, Any]], meta: dict[str, Any]) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    blank = prs.slide_layouts[6]
    rgb = lambda c: RGBColor.from_string(c)  # noqa: E731
    lb = logo_bytes()
    FONT = "Arial"

    def text(s, x, y, w, h, runs, size=7, color=C["text"], bold=False, align=None, wrap=True):
        tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(1)
        for i, item in enumerate(runs if isinstance(runs, list) else [runs]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if isinstance(item, tuple):
                t, sz, col, b = item
            else:
                t, sz, col, b = item, size, color, bold
            r = p.add_run()
            r.text = plain(t)
            r.font.size, r.font.bold, r.font.name = Pt(sz), b, FONT
            r.font.color.rgb = rgb(col)
            if align:
                p.alignment = align
        return tb

    def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
        r = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
        r.shadow.inherit = False
        if fill:
            r.fill.solid()
            r.fill.fore_color.rgb = rgb(fill)
        else:
            r.fill.background()
        if line:
            r.line.color.rgb = rgb(line)
            r.line.width = Pt(.6)
        else:
            r.line.fill.background()
        return r

    def header(s, page):
        if lb:
            s.shapes.add_picture(io.BytesIO(lb), Pt(22), Pt(14), height=Pt(30))
        long_title = len(page["title"]) * 11.5 * 0.52 > 540
        tsize = 9.6 if long_title else 11.5
        text(s, 128, 8, 545, 28, page["title"], tsize, C["navy"], True)
        text(s, 128, 36, 545, 12, page["subtitle"], 7.0, C["orange"], True)
        tag = plain(page["tag"])
        tw = len(tag) * 3.7 + 14
        pill = rect(s, 938 - tw, 6, tw, 11, C["navy"])
        pill.text_frame.text = tag
        pill.text_frame.word_wrap = False
        pill.text_frame.margin_left = pill.text_frame.margin_right = Pt(3)
        pill.text_frame.margin_top = pill.text_frame.margin_bottom = Pt(0)
        for p_ in pill.text_frame.paragraphs:
            p_.alignment = PP_ALIGN.CENTER
            for r_ in p_.runs:
                r_.font.size, r_.font.bold, r_.font.name = Pt(6), True, FONT
                r_.font.color.rgb = rgb(C["white"])
        text(s, 690, 18, 248, 10, meta["period_text"], 6, C["muted"], align=PP_ALIGN.RIGHT)
        if page.get("zone_label"):
            text(s, 690, 26, 248, 10, page["zone_label"], 6.2, C["navy"], True, align=PP_ALIGN.RIGHT)
        badge = rect(s, 778, 37, 160, 12, C["navy"])
        badge.text_frame.text = "CONFIDENTIAL · C-SUITE ANALYTICS"
        for p in badge.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size, r.font.bold, r.font.name = Pt(6.5), True, FONT
                r.font.color.rgb = rgb(C["white"])
        line = s.shapes.add_connector(1, Pt(22), Pt(52), Pt(938), Pt(52))
        line.line.color.rgb = rgb(C["navy"])
        line.line.width = Pt(1.2)

    def section(s, title, x, y, w, color=C["navy"]):
        main, accent = sec_text(title)
        full_len = len(main) + (len(accent) + 3 if accent else 0)
        size = min(7.0, max(5.0, w / (full_len * 0.62))) if full_len else 7.0
        runs = [(main.upper(), size, color, True)]
        tb = text(s, x, y, w, 10, runs, size, color, True)
        if accent:
            r = tb.text_frame.paragraphs[0].add_run()
            r.text = "  · " + accent.upper()
            r.font.size, r.font.bold, r.font.name = Pt(size), True, FONT
            r.font.color.rgb = rgb(C["orange"])
        ln = s.shapes.add_connector(1, Pt(x), Pt(y + 11), Pt(x + w), Pt(y + 11))
        ln.line.color.rgb = rgb(C["line"])
        ln.line.width = Pt(.5)
        return y + 14

    def table(s, tbl, x, y, w, max_h):
        from pptx.enum.text import MSO_ANCHOR

        cols = tbl["columns"]
        row_h = 11.5
        n_head = 2 if tbl.get("groups") else 1
        max_h -= 14
        max_rows = max(3, int(max_h / row_h) - n_head)
        rows = tbl["rows"][:max_rows]
        shape = s.shapes.add_table(len(rows) + n_head, len(cols), Pt(x), Pt(y), Pt(w), Pt(row_h * (len(rows) + n_head)))
        t = shape.table
        first = w * (0.30 if len(cols) <= 10 else 0.22)
        t.columns[0].width = Pt(first)
        for ci in range(1, len(cols)):
            t.columns[ci].width = Pt((w - first) / (len(cols) - 1))
        size = 6 if len(cols) <= 9 else 5.5

        def fill(cell, value, color=C["text"], bg=None, bold=False, align=PP_ALIGN.RIGHT, italic=False, dot=None):
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Pt(2)
            cell.margin_top = cell.margin_bottom = Pt(0)
            tf = cell.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.alignment = align
            if dot:
                rd = p.add_run()
                rd.text = "● "
                rd.font.size, rd.font.name = Pt(size), FONT
                rd.font.color.rgb = rgb(dot)
            r = p.add_run()
            r.text = plain(value)
            r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, FONT
            r.font.color.rgb = rgb(color)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(bg or C["white"])

        if tbl.get("groups"):
            col = 0
            for label, span in tbl["groups"]:
                if span > 1:
                    t.cell(0, col).merge(t.cell(0, col + span - 1))
                fill(t.cell(0, col), label, C["white"], C["navy"], True, PP_ALIGN.LEFT if col == 0 else PP_ALIGN.CENTER)
                col += span
            for ci, c in enumerate(cols):
                fill(t.cell(1, ci), c, C["navy"], "EDF0F4", True, PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT)
        else:
            for ci, c in enumerate(cols):
                fill(t.cell(0, ci), c, C["white"], C["navy"], True, PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT)
        for i in range(n_head):
            t.rows[i].height = Pt(row_h)
        for ri, r in enumerate(rows, start=n_head):
            t.rows[ri].height = Pt(row_h)
            if r["kind"] == "group":
                t.cell(ri, 0).merge(t.cell(ri, len(cols) - 1))
                fill(t.cell(ri, 0), r["cells"][0], C["white"], C["navy"], True, PP_ALIGN.LEFT)
                continue
            bg = C["hero"] if r["kind"] == "hero" else None
            for ci, (v, tone) in enumerate(zip(r["cells"], r["tones"])):
                if tone and tone.startswith("status:"):
                    fill(t.cell(ri, ci), v, C["text"], bg, True, PP_ALIGN.LEFT, dot=TONE.get(tone.split(":")[1], C["muted"]))
                elif ci == 0:
                    fill(t.cell(ri, ci), v, "9A6A00" if r["kind"] == "hero" else C["text"], bg, r["kind"] != "hero",
                         PP_ALIGN.LEFT, italic=r["kind"] == "hero")
                else:
                    fill(t.cell(ri, ci), v, TONE.get(tone, C["text"]), bg, tone in TONE)
        end = y + row_h * (len(rows) + n_head) + 3 * n_head
        if tbl.get("note"):
            text(s, x, end + 4, w, 10, tbl["note"], 5.5, C["muted"])
        return end + 14

    def chart(s, ch, x, y, w, h):
        text(s, x, y, w, 10, f"{ch['judul']} ({ch['satuan']})", 6.5, C["navy"], True)
        data = CategoryChartData()
        data.categories = [_wrap(c, 14, 2) for c in ch["kategori"]]
        for se in ch["seri"]:
            data.add_series(se["nama"], se["nilai"])
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(x), Pt(y + 10), Pt(w), Pt(h - 10), data)
        c = gf.chart
        c.has_title = False
        c.has_legend = True
        c.legend.position, c.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        c.legend.font.size = Pt(6)
        c.font.size, c.font.name = Pt(6), FONT
        c.value_axis.tick_labels.number_format, c.value_axis.tick_labels.number_format_is_linked = '#,##0', False
        vals = [v for se in ch["seri"] for v in se["nilai"] if v is not None]
        if vals:
            ticks = nice_ticks(min(0.0, min(vals)), max(0.0, max(vals)), 4)
            c.value_axis.minimum_scale, c.value_axis.maximum_scale = ticks[0], ticks[-1]
            if len(ticks) > 1:
                c.value_axis.major_unit = ticks[1] - ticks[0]
        c.value_axis.major_gridlines.format.line.color.rgb = rgb(C["line"])
        c.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        plot = c.plots[0]
        plot.gap_width, plot.overlap = 50, -5
        for i, se in enumerate(plot.series):
            se.format.fill.solid()
            se.format.fill.fore_color.rgb = rgb(series_colors(ch)[i % len(ch["seri"])])
            se.invert_if_negative = False

    def card_block(s, items, x, y, w, y_max):
        for it in items:
            body = it.get("uraian", "")
            lines = 2 + math.ceil(len(body) * 3.2 / max(w, 1))
            h = 8 + lines * 8
            if y + h > y_max:
                break
            rect(s, x, y, w, h, "F7F9FB")
            rect(s, x, y, 3, h, LEVEL.get(it.get("tingkat"), C["accent"]))
            text(s, x + 6, y + 2, w - 8, h - 2, [(it.get("kategori", ""), 5.5, C["muted"], True),
                                                 (it.get("judul", ""), 7, C["navy"], True), (body, 6.2, C["text"], False)])
            y += h + 4
        return y

    for page in pages:
        s = prs.slides.add_slide(blank)
        header(s, page)
        top = 60
        if page["kind"] == "exec":
            cw = (960 - 44 - 24) / 4
            for i, z in enumerate(page["zone_cards"]):
                x = 22 + i * (cw + 8)
                card_h = 24 + 11 * len(z["metrics"]) + (14 if z["note"] else 0)
                rect(s, x, top, cw, card_h, None, C["line"])
                rect(s, x, top, cw, 3, list(ZONE_COLORS.values())[i])
                text(s, x + 5, top + 5, cw - 80, 10, z["title"], 6.8, C["navy"], True)
                nick = rect(s, x + cw - 74, top + 6, 68, 10, C["navy"])
                nick.text_frame.text = z["subtitle"]
                for p_ in nick.text_frame.paragraphs:
                    p_.alignment = PP_ALIGN.CENTER
                    for r_ in p_.runs:
                        r_.font.size, r_.font.bold, r_.font.name = Pt(5.5), True, FONT
                        r_.font.color.rgb = rgb(C["white"])
                yy = top + 18
                for m in z["metrics"]:
                    text(s, x + 5, yy, cw * .62, 10, m["label"], 6)
                    text(s, x + cw * .6, yy, cw * .38 - 5, 10, m["value"], 6.5, TONE.get(m["tone"], C["text"]), True,
                         align=PP_ALIGN.RIGHT)
                    yy += 11
                if z["note"]:
                    text(s, x + 5, yy + 1, cw - 10, 14, f"{z['note_label']} {z['note']}", 5.6, C["muted"])
            y = top + 24 + 11 * max(len(z["metrics"]) for z in page["zone_cards"]) + 22
            if page["bullets"]:
                text(s, 22, y, 916, 40, ["• " + b for b in page["bullets"][:4]], 6.6)
                y += 10 * min(len(page["bullets"]), 4) + 6
            lw = 916 * .58
            ty = section(s, page["sec"]["table"], 22, y, lw)
            table(s, page["table"], 22, ty, lw, 520 - ty)
            rx, rw = 22 + lw + 14, 916 - lw - 14
            ry = section(s, page["sec"]["insights"], rx, y, rw)
            ry = card_block(s, page["insights"], rx, ry, rw, 390)
            cw2 = (rw - 8) / 2
            ch_h = min(170, 518 - ry - 4)
            for i, ch in enumerate(page["charts"][:2]):
                chart(s, ch, rx + i * (cw2 + 8), ry + 4, cw2, ch_h)
        else:
            kw = (960 - 44 - 24) / 4
            for i, k in enumerate(page["kpis"]):
                x = 22 + i * (kw + 8)
                rect(s, x, top, kw, 58, None, C["line"])
                rect(s, x, top, kw, 3, TONE.get(k["chip_tone"], C["accent"]))
                text(s, x + 5, top + 4, kw - 54, 10, k["label"], 5.8, C["muted"], True)
                text(s, x + 5, top + 13, kw - 10, 16, k["value"], 12, TONE.get(k["tone"], C["text"]), True)
                text(s, x + 5, top + 29, kw - 10, 9, k["detail"], 5.6, C["mid"], True)
                text(s, x + 5, top + 37, kw - 10, 9, k["extra"], 5.3, C["muted"])
                chip = rect(s, x + 5, top + 46, 62, 8, {"low": "E8F5EC", "mid": "FFF4D6", "high": "FDECEA"}.get(k["chip_tone"], C["light"]),
                            TONE.get(k["chip_tone"], C["muted"]))
                chip.text_frame.text = plain(k["chip"])
                chip.text_frame.margin_top = chip.text_frame.margin_bottom = Pt(0)
                for p in chip.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
                    for r in p.runs:
                        r.font.size, r.font.bold, r.font.name = Pt(5.2), True, FONT
                        r.font.color.rgb = rgb(TONE.get(k["chip_tone"], C["muted"]))
            y = top + 64
            w1, w2 = 916 * .45, 916 * .28
            w3 = 916 - w1 - w2 - 20
            x1, x2, x3 = 22, 22 + w1 + 10, 22 + w1 + w2 + 20
            ty = section(s, page["sec"]["table"], x1, y, w1)
            table(s, page["table"], x1, ty, w1, 512 - ty)
            cy = section(s, page["sec"]["charts"], x2, y, w2)
            n_ch = max(len(page["charts"]), 1)
            anom_h = 16 + 10 * len(page["anomaly"]) if page["anomaly"] else 0
            ch_h = min(118, (520 - cy - anom_h) / n_ch)
            for ch in page["charts"]:
                chart(s, ch, x2, cy, w2, ch_h)
                cy += ch_h + 2
            if page["anomaly"]:
                rect(s, x2, cy, w2, anom_h, C["anom"], C["gold"])
                text(s, x2 + 4, cy + 2, w2 - 8, anom_h - 2, [(page["anomaly_title"], 6.3, C["mid"], True)]
                     + [("• " + a, 6, C["text"], False) for a in page["anomaly"]])
            iy = section(s, page["sec"]["insights"], x3, y, w3)
            iy = card_block(s, page["insights"], x3, iy, w3, 300)
            if page["watch"]:
                iy = section(s, page["sec"]["watch"], x3, iy + 2, w3, C["high"])
                for wv in page["watch"]:
                    if iy + 52 > 514:
                        break
                    rect(s, x3, iy, w3, 50, "F4FBF6", C["line"])
                    text(s, x3 + 4, iy + 1, w3 - 44, 9, wv["product"], 6.2, C["navy"], True)
                    if wv["brand"]:
                        b_ = rect(s, x3 + w3 - 38, iy + 2, 34, 8, C["high"])
                        b_.text_frame.text = wv["brand"]
                        b_.text_frame.margin_top = b_.text_frame.margin_bottom = Pt(0)
                        for p_ in b_.text_frame.paragraphs:
                            p_.alignment = PP_ALIGN.CENTER
                            for r_ in p_.runs:
                                r_.font.size, r_.font.bold, r_.font.name = Pt(5), True, FONT
                                r_.font.color.rgb = rgb(C["white"])
                    text(s, x3 + 4, iy + 9, w3 - 8, 8, wv.get("subtitle") or f"{wv['segment']} · {wv['metric']} = PTPL − produk ini", 5.1, C["muted"])
                    sw = (w3 - 8) / max(len(wv["stats"]), 1)
                    for si, (lbl, v, t) in enumerate(wv["stats"]):
                        rect(s, x3 + 4 + si * sw, iy + 17, sw - 2, 13, C["light"])
                        text(s, x3 + 4 + si * sw, iy + 17, sw - 2, 13, [(lbl, 4.6, C["muted"], False), (v, 5.6, TONE.get(t, C["text"]), True)],
                             align=PP_ALIGN.CENTER)
                    mw = (w3 - 8) / max(len(wv["months"]), 1)
                    for mi, (m, v, t) in enumerate(zip(wv["months"], wv["values"], wv["tones"])):
                        bgc = {"low": "E8F5EC", "high": "FDECEA"}.get(t)
                        if bgc:
                            rect(s, x3 + 4 + mi * mw, iy + 32, mw - 2, 14, bgc)
                        text(s, x3 + 4 + mi * mw, iy + 32, mw - 2, 14, [(m, 4.6, C["muted"], False), (v, 5.4, TONE.get(t, C["text"]), True)],
                             align=PP_ALIGN.CENTER)
                    iy += 53
            if page.get("footnote"):
                text(s, 22, 516, 900, 9, page["footnote"], 5.4, C["muted"])
        text(s, 22, 526, 880, 10, page["footer"], 5.6, C["muted"])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

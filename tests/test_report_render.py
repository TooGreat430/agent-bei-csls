"""Uji render laporan ke HTML, PDF, dan PPTX (tanpa GCP)."""
import io
import json
import os
import sys
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tests"))
import _stubs  # noqa: E402

_stubs.install()
os.environ["LIB_TEMPLATE_SOURCE"] = "local"

from library_agent import report_engine, report_render as rr  # noqa: E402

try:
    import reportlab  # noqa: F401
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False
try:
    import pptx  # noqa: F401
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

META = {"company_name": "PT Pertamina Lubricants", "report_title": "Laporan Uji", "generated_date": "30-09-2026"}


def load(name):
    with open(os.path.join(ROOT, "tests", "fixtures", name), encoding="utf-8") as fh:
        return json.load(fh)


class LayoutTest(unittest.TestCase):
    def test_default_layout_from_schema(self):
        t = report_engine.load_template("laporan_studi")
        kinds = [(b["type"], b.get("field")) for b in rr.resolve_layout(t.manifest, t.schema)]
        self.assertIn(("findings", "temuan_utama"), kinds)
        self.assertIn(("bullets", "implikasi"), kinds)
        self.assertIn(("table", "rekomendasi"), kinds)
        self.assertNotIn(("text", "periode"), kinds)  # periode dipakai di sampul

    def test_chunking(self):
        chunks = rr.chunk_by_lines(["x" * 200] * 6, chars_per_line=100, max_lines=6)
        self.assertEqual([len(c) for c in chunks], [2, 2, 2])
        rows = rr.chunk_rows([["a", "b" * 150]] * 5, [10, 50], max_lines=7)
        self.assertEqual([len(c) for c in rows], [2, 2, 1])

    def test_formats(self):
        self.assertEqual(report_engine.normalize_format("PowerPoint"), "pptx")
        self.assertEqual(report_engine.normalize_format(".PDF"), "pdf")
        t = report_engine.load_template("studi_bei_nps")
        self.assertEqual(t.outputs, ["html", "pdf", "pptx"])
        with self.assertRaises(ValueError):
            report_engine.render(t, {}, META, "docx")

    def test_finding_table_pads_rows(self):
        cols, rows = rr.finding_table({"tabel": {"kolom": ["A", "B", "C"], "baris": [["1"], ["1", "2", "3", "4"]]}})
        self.assertEqual(rows, [["1", "", ""], ["1", "2", "3"]])
        self.assertIsNone(rr.finding_table({"tabel": {"kolom": [], "baris": []}}))


class RenderTest(unittest.TestCase):
    def setUp(self):
        self.t = report_engine.load_template("studi_bei_nps")
        self.content = load("sample_bei_nps.json")

    def test_html(self):
        data, ext, _ = report_engine.render(self.t, self.content, META, "html")
        html = data.decode()
        self.assertEqual(ext, "html")
        self.assertIn("BRAND EQUITY INDEX".lower(), html.lower())
        self.assertIn("[Studi BEI &amp; NPS Motor Q1 2026, v1, hal. 31]", html)

    @unittest.skipUnless(HAS_REPORTLAB, "reportlab belum terpasang")
    def test_pdf(self):
        data, ext, ctype = report_engine.render(self.t, self.content, META, "pdf")
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertEqual((ext, ctype), ("pdf", "application/pdf"))

    @unittest.skipUnless(HAS_PPTX, "python-pptx belum terpasang")
    def test_pptx_and_long_content_splits(self):
        from pptx import Presentation

        content = json.loads(json.dumps(self.content))
        content["ringkasan_eksekutif"] = ["Poin panjang " * 30] * 6
        data, ext, _ = report_engine.render(self.t, content, META, "pptx")
        prs = Presentation(io.BytesIO(data))
        titles = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
        self.assertTrue(any("RINGKASAN EKSEKUTIF (LANJUTAN)" in t for t in titles))
        self.assertEqual(ext, "pptx")

    @unittest.skipUnless(HAS_PPTX and HAS_REPORTLAB, "library render belum terpasang")
    def test_legacy_template_all_formats(self):
        t = report_engine.load_template("laporan_studi")
        content = load("sample_content.json")
        meta = report_engine.build_meta(t, "penguji@klien.co.id", "Laporan Uji")
        for fmt in ("html", "pdf", "pptx"):
            data, _, _ = report_engine.render(t, content, meta, fmt)
            self.assertGreater(len(data), 1000)


if __name__ == "__main__":
    unittest.main()

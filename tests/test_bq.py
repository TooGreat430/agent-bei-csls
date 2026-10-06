"""Uji integrasi data BigQuery, pengaman angka, grafik, dan routing (tanpa GCP)."""
import json
import os
import sys
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tests"))
import _stubs  # noqa: E402

_stubs.install()
os.environ["LIB_TEMPLATE_SOURCE"] = "local"

from library_agent import data_agent, grounding, report_engine  # noqa: E402
from library_agent import report_render as rr  # noqa: E402

try:
    import pptx  # noqa: F401
    import reportlab  # noqa: F401
    HAS_RENDER = True
except ImportError:
    HAS_RENDER = False


def fixture(name):
    with open(os.path.join(ROOT, "tests", "fixtures", name), encoding="utf-8") as fh:
        return json.load(fh)


class DataAgentParseTest(unittest.TestCase):
    def test_parse_stream(self):
        stream = [
            {"system_message": {"text": {"parts": ["Sedang menganalisis..."], "text_type": "PROGRESS"}}},
            {"system_message": {"data": {"generated_sql": "SELECT 1"}}},
            {"system_message": {"data": {"result": {
                "schema": {"fields": [{"name": "ZONE"}, {"name": "GAP_HJ"}]},
                "data": [{"ZONE": "Zona 1", "GAP_HJ": -9656.4}, {"ZONE": "Zona 2", "GAP_HJ": float("nan")}]}}}},
            {"system_message": {"text": {"parts": ["Gap HJ Zona 1 −Rp 9.656/L."], "text_type": "FINAL_RESPONSE"}}},
            {"system_message": {"chart": {"query": {}}}},
        ]
        out = data_agent.parse_stream(stream)
        self.assertEqual(out["answer"], "Gap HJ Zona 1 −Rp 9.656/L.")
        self.assertEqual(out["tables"][0]["columns"], ["ZONE", "GAP_HJ"])
        self.assertEqual(out["tables"][0]["rows"], [["Zona 1", -9656.4], ["Zona 2", None]])
        self.assertEqual(out["sql"], ["SELECT 1"])
        md = data_agent.table_to_markdown(out["tables"][0])
        self.assertIn("| ZONE | GAP_HJ |", md)

    def test_errors_collected(self):
        out = data_agent.parse_stream([{"system_message": {"error": {"text": "permission denied"}}}])
        self.assertEqual(out["errors"], ["permission denied"])
        self.assertEqual(out["answer"], "")


class UserAuthTest(unittest.TestCase):
    def test_token_lookup_plain_and_temp(self):
        self.assertEqual(data_agent.find_user_token({"mia-bigquery": "tok1"}, "mia-bigquery"), "tok1")
        self.assertEqual(data_agent.find_user_token({"temp:mia-bigquery": "tok2"}, "mia-bigquery"), "tok2")
        self.assertIsNone(data_agent.find_user_token({}, "mia-bigquery"))
        self.assertIsNone(data_agent.find_user_token({"mia-bigquery": ""}, "mia-bigquery"))

    def test_error_classification(self):
        self.assertEqual(data_agent.classify_error("401 Request had invalid authentication"), "token")
        self.assertEqual(data_agent.classify_error("403 Permission denied on table"), "permission")
        self.assertEqual(data_agent.classify_error("500 internal"), "other")


class NumbersTest(unittest.TestCase):
    def test_parse_number_formats(self):
        cases = {"−9.390/L": -9390, "Rp 10.596/L": 10596, "+666/L": 666, "1,3": 1.3, "16.5": 16.5,
                 "−12.592": -12592, "-4,5%": -4.5, "1.234.567": 1234567, "NA": None, "": None}
        for raw, expected in cases.items():
            self.assertEqual(rr.parse_number(raw), expected, raw)

    def test_value_tone_and_status(self):
        self.assertEqual(rr.value_tone("−9.390", "negatif"), "low")
        self.assertEqual(rr.value_tone("+666", "negatif"), "high")
        self.assertEqual(rr.value_tone("+963", "positif"), "low")
        self.assertIsNone(rr.value_tone("AMAN", "netral"))
        self.assertEqual(rr.status_tone("kritis"), "high")
        self.assertEqual(rr.fmt_id(-9390), "−9.390")

    def test_clean_charts_drops_misaligned(self):
        charts = rr.clean_charts([
            {"judul": "ok", "jenis": "bar", "satuan": "IDR/L", "kategori": ["a", "b"],
             "seri": [{"nama": "s1", "nilai": [1, "−2.000"]}, {"nama": "s2", "nilai": [1]}]},
            {"judul": "kosong", "kategori": ["a"], "seri": [{"nama": "x", "nilai": [None]}]},
        ])
        self.assertEqual(len(charts), 1)
        self.assertEqual(charts[0]["seri"], [{"nama": "s1", "nilai": [1.0, -2000.0]}])


class GroundingTest(unittest.TestCase):
    def setUp(self):
        self.insights = [{"content": "Gap HJ nasional Matic-S −Rp 9.390/L.", "citations": [],
                          "data": {"columns": ["ZONE", "GAP"], "rows": [["Zona 1", -9656.4], ["Zona 2", -9487.2]]}}]
        self.nums = grounding.source_numbers(self.insights)

    def test_grounded_values_pass(self):
        content = {"grafik": [{"seri": [{"nilai": [-9390, -9656, -9487]}]}],
                   "kpi_zona": [{"metrik": [{"label": "x", "nilai": "−9.656/L"}]}]}
        self.assertEqual(grounding.check(content, self.nums), [])

    def test_invented_values_flagged_and_stripped(self):
        content = {"grafik": [{"seri": [{"nilai": [-9390, -12345]}]}],
                   "matriks": {"baris": [{"label": "Zona 9", "nilai": ["−77.777", "AMAN", "NA"]}]},
                   "kpi_zona": [{"judul": "A", "metrik": [{"label": "x", "nilai": "−55.555/L"}]}]}
        issues = grounding.check(content, self.nums)
        self.assertEqual(len(issues), 3)
        fixed, removed = grounding.strip_ungrounded(content, self.nums)
        self.assertEqual(removed, 3)
        self.assertEqual(fixed["grafik"][0]["seri"][0]["nilai"], [-9390, None])
        self.assertEqual(fixed["matriks"]["baris"][0]["nilai"], ["NA", "AMAN", "NA"])
        self.assertEqual(fixed["kpi_zona"], [])


class TemplateTest(unittest.TestCase):
    def test_bq_template_metadata(self):
        t = report_engine.load_template("daya_saing_harga_bq")
        self.assertTrue(t.manifest["grounding"])
        self.assertEqual(t.outputs, ["html", "pdf", "pptx"])

    @unittest.skipUnless(HAS_RENDER, "library render belum terpasang")
    def test_bq_template_renders_all_formats(self):
        from pptx import Presentation
        import io

        t = report_engine.load_template("daya_saing_harga_bq")
        content = fixture("sample_bq_report.json")
        meta = report_engine.build_meta(t, "u", "Uji BQ")
        html, _, _ = report_engine.render(t, content, meta, "html")
        self.assertIn(b"<svg", html)
        self.assertIn(b"KRITIS", html)
        pdf, _, _ = report_engine.render(t, content, meta, "pdf")
        self.assertTrue(pdf.startswith(b"%PDF"))
        data, _, _ = report_engine.render(t, content, meta, "pptx")
        prs = Presentation(io.BytesIO(data))
        charts = [sh for s in prs.slides for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart]
        self.assertEqual(len(charts), 2)


class SourceRequirementTest(unittest.TestCase):
    def test_bq_template_rejects_document_only_insights(self):
        t = report_engine.load_template("daya_saing_harga_bq")
        doc_only = [{"source": "dokumen", "content": "x"}]
        self.assertIn("data BigQuery", report_engine.source_requirement_problem(t, doc_only))
        mixed = doc_only + [{"source": "bigquery", "content": "y"}]
        self.assertIsNone(report_engine.source_requirement_problem(t, mixed))

    def test_document_templates_accept_any_source(self):
        t = report_engine.load_template("studi_bei_nps")
        self.assertIsNone(report_engine.source_requirement_problem(t, [{"source": "bigquery"}]))

    def test_literal_newlines_become_paragraphs(self):
        self.assertEqual(rr.paragraphs("Satu.\\n\\nDua."), ["Satu.", "Dua."])


class PreviewTest(unittest.TestCase):
    def test_every_template_has_sample(self):
        for tid in report_engine.list_template_ids():
            self.assertIsNotNone(report_engine.load_sample(tid), tid)

    @unittest.skipUnless(HAS_RENDER, "library render belum terpasang")
    def test_preview_is_marked(self):
        t = report_engine.load_template("daya_saing_harga_bq")
        meta = report_engine.preview_meta(t, "u")
        html, _, _ = report_engine.render(t, report_engine.load_sample("daya_saing_harga_bq"), meta, "html")
        self.assertIn("CONTOH TAMPILAN".encode(), html)

    def test_manifest_version_compare(self):
        sys.path.insert(0, os.path.join(ROOT, "scripts"))
        import importlib.util
        spec = importlib.util.spec_from_file_location("ut", os.path.join(ROOT, "scripts", "upload_templates.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        self.assertGreater(mod._version("2"), mod._version("1"))
        self.assertGreater(mod._version("1.10"), mod._version("1.9"))


class RoutingTest(unittest.TestCase):
    def test_every_turn_starts_at_root(self):
        from library_agent.subagents.data import data_agent as d
        from library_agent.subagents.report import report_agent as r
        from library_agent.subagents.research import research_agent as s

        for agent in (d, r, s):
            self.assertTrue(agent.disallow_transfer_to_parent, agent.name)


if __name__ == "__main__":
    unittest.main()


class ModelSchemaTest(unittest.TestCase):
    def test_bq_schema_simplified_for_model(self):
        t = report_engine.load_template("daya_saing_harga_bq")
        text = json.dumps(report_engine.schema_for_model(t.schema))
        for key in ("anyOf", "$ref", "additionalProperties", "minItems", "maxItems", "$schema"):
            self.assertNotIn(key, text, key)
        self.assertIn('"enum"', text)  # enum tetap ada

    def test_refs_inlined(self):
        schema = {"type": "object", "properties": {"a": {"$ref": "#/$defs/x"}},
                  "$defs": {"x": {"type": "object", "properties": {"b": {"type": ["string", "null"]}}}}}
        out = report_engine.schema_for_model(schema)
        self.assertEqual(out["properties"]["a"]["properties"]["b"], {"type": "string"})
        self.assertNotIn("$defs", out)

    def test_report_tool_never_raises(self):
        from library_agent.tools import report_tools

        @report_tools._logged
        def boom(tool_context=None):
            raise RuntimeError("400 INVALID_ARGUMENT")

        result = boom()
        self.assertEqual(result["status"], "error")
        self.assertIn("kendala teknis", result["message"])
